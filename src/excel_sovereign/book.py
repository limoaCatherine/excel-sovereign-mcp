"""openpyxl writes. One load, one save, keep_vba. This module is the only openpyxl caller."""

from __future__ import annotations

import os
import shutil
import tempfile
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from excel_sovereign.lock import canonical_path


@dataclass
class BookResult:
    wrote: list[dict] = field(default_factory=list)
    formula_cells: list[dict] = field(default_factory=list)
    completed: int = 0
    failed_at: int | None = None
    error: str | None = None
    touched_sheets: list[str] = field(default_factory=list)


def load_for_edit(path: str):
    return load_workbook(path, keep_vba=path.lower().endswith(".xlsm"))


def create_empty(path: str) -> None:
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    workbook.save(path)
    workbook.close()


def save_atomic(workbook, path: str) -> str:
    """Save to a sibling temp file. Caller replaces the target after verification prep."""
    directory = os.path.dirname(path) or "."
    suffix = ".xlsm" if path.lower().endswith(".xlsm") else ".xlsx"
    fd, temp_path = tempfile.mkstemp(prefix=".excel-sovereign-mcp-", suffix=suffix, dir=directory)
    os.close(fd)
    try:
        workbook.save(temp_path)
    except Exception:
        try:
            os.remove(temp_path)
        except OSError:
            pass
        raise
    return temp_path


def replace_target(temp_path: str, path: str) -> None:
    os.replace(temp_path, path)


def apply_ops(workbook, ops: list[dict]) -> BookResult:
    result = BookResult()
    for index, op in enumerate(ops):
        try:
            spans = _apply_one(workbook, op)
        except Exception as exc:  # operation error before save
            result.failed_at = index
            result.error = str(exc)
            return result
        result.completed = index + 1
        for sheet, address, formulas in spans:
            if sheet not in result.touched_sheets:
                result.touched_sheets.append(sheet)
            if address:
                result.wrote.append({"sheet": sheet, "range": address})
            for cell, formula in formulas:
                result.formula_cells.append({"sheet": sheet, "cell": cell, "formula": formula})
    return result


def _sheet(workbook, name: str) -> Worksheet:
    if name not in workbook.sheetnames:
        raise KeyError(f"sheet not found: {name}")
    return workbook[name]


def _apply_one(workbook, op: dict) -> list[tuple[str, str | None, list[tuple[str, str]]]]:
    action = op["action"]
    if action == "create_workbook":
        return []
    if action == "create_sheet":
        name = str(op.get("name") or op.get("sheet"))
        if name in workbook.sheetnames:
            raise ValueError(f"sheet already exists: {name}")
        workbook.create_sheet(name)
        return [(name, "A1", [])]
    if action == "set_values":
        return [_write_block(workbook, op, formulas=False)]
    if action == "set_formulas":
        return [_write_block(workbook, op, formulas=True)]
    if action == "define_name":
        _define_name(workbook, op, replace=False)
        return []
    if action == "define_name_update":
        _define_name(workbook, op, replace=True)
        return []
    if action == "format":
        _format(workbook, op)
        return [(str(op["sheet"]), str(op["range"]), [])]
    if action == "merge":
        _sheet(workbook, op["sheet"]).merge_cells(str(op["range"]))
        return [(str(op["sheet"]), str(op["range"]), [])]
    if action == "unmerge":
        _sheet(workbook, op["sheet"]).unmerge_cells(str(op["range"]))
        return [(str(op["sheet"]), str(op["range"]), [])]
    if action == "set_row_height":
        ws = _sheet(workbook, op["sheet"])
        row = int(op["row"])
        count = int(op.get("count") or 1)
        height = float(op["height"])
        for offset in range(count):
            ws.row_dimensions[row + offset].height = height
        return [(str(op["sheet"]), f"A{row}:A{row + count - 1}", [])]
    if action == "set_column_width":
        ws = _sheet(workbook, op["sheet"])
        column = str(op["column"])
        ws.column_dimensions[column].width = float(op["width"])
        return [(str(op["sheet"]), f"{column}1", [])]
    if action == "clear_contents":
        _clear(workbook, op, formats=False)
        return [(str(op["sheet"]), str(op["range"]), [])]
    if action == "clear_all":
        _clear(workbook, op, formats=True)
        return [(str(op["sheet"]), str(op["range"]), [])]
    raise ValueError(f"openpyxl cannot apply {action}")


def _write_block(workbook, op: dict, formulas: bool):
    ws = _sheet(workbook, op["sheet"])
    rows = op["formulas"] if formulas else op["values"]
    if not isinstance(rows, list) or not rows:
        raise ValueError("values or formulas must be a non-empty 2D array")
    min_col, min_row, _, _ = range_boundaries(str(op["range"]))
    width = max(len(row) for row in rows)
    _reject_merged_write(ws, min_row, min_col, min_row + len(rows) - 1, min_col + width - 1)
    found: list[tuple[str, str]] = []
    for r_index, row in enumerate(rows):
        if not isinstance(row, list):
            raise ValueError("each row must be a list")
        for c_index in range(width):
            value = row[c_index] if c_index < len(row) else None
            cell = ws.cell(min_row + r_index, min_col + c_index, value)
            if isinstance(value, str) and value.startswith("="):
                found.append((cell.coordinate, value))
    end = ws.cell(min_row + len(rows) - 1, min_col + width - 1).coordinate
    start = ws.cell(min_row, min_col).coordinate
    address = start if start == end else f"{start}:{end}"
    return (str(op["sheet"]), address, found)


def _reject_merged_write(ws: Worksheet, min_row: int, min_col: int, max_row: int, max_col: int) -> None:
    for merged in ws.merged_cells.ranges:
        overlaps = not (
            merged.max_row < min_row
            or merged.min_row > max_row
            or merged.max_col < min_col
            or merged.min_col > max_col
        )
        if not overlaps:
            continue
        only_top_left = (
            min_row == max_row == merged.min_row and min_col == max_col == merged.min_col
        )
        if not only_top_left:
            raise ValueError(f"write intersects merged range {merged.coord}")


def _define_name(workbook, op: dict, replace: bool) -> None:
    name = str(op["name"])
    formula = str(op.get("formula") or op.get("reference") or "")
    if formula.startswith("="):
        formula = formula[1:]
    existing = name in workbook.defined_names
    if existing and not replace:
        raise ValueError(f"name already exists: {name}")
    if existing:
        del workbook.defined_names[name]
    workbook.defined_names.add(DefinedName(name=name, attr_text=formula))


def _format(workbook, op: dict) -> None:
    ws = _sheet(workbook, op["sheet"])
    min_col, min_row, max_col, max_row = range_boundaries(str(op["range"]))
    font_kwargs = {}
    if "bold" in op:
        font_kwargs["bold"] = bool(op["bold"])
    if "italic" in op:
        font_kwargs["italic"] = bool(op["italic"])
    if "underline" in op:
        font_kwargs["underline"] = "single" if op["underline"] else None
    if op.get("fontSize") is not None:
        font_kwargs["size"] = op["fontSize"]
    if op.get("fontName"):
        font_kwargs["name"] = op["fontName"]
    if op.get("fontColor"):
        font_kwargs["color"] = _hex(op["fontColor"])
    fill = None
    if op.get("fillColor"):
        fill = PatternFill("solid", fgColor=_hex(op["fillColor"]))
    border = None
    if op.get("borderStyle"):
        side = Side(style=_border_style(op["borderStyle"]), color=_hex(op.get("borderColor") or "000000"))
        if op.get("borderEdges") == "top":
            border = Border(top=side)
        else:
            border = Border(left=side, right=side, top=side, bottom=side)
    alignment_kwargs = {}
    if op.get("horizontalAlignment"):
        alignment_kwargs["horizontal"] = op["horizontalAlignment"]
    if op.get("verticalAlignment"):
        alignment_kwargs["vertical"] = "center" if op["verticalAlignment"] == "middle" else op["verticalAlignment"]
    if "wrapText" in op:
        alignment_kwargs["wrap_text"] = bool(op["wrapText"])
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if font_kwargs:
                cell.font = Font(
                    name=font_kwargs.get("name", cell.font.name),
                    size=font_kwargs.get("size", cell.font.size),
                    bold=font_kwargs.get("bold", cell.font.bold),
                    italic=font_kwargs.get("italic", cell.font.italic),
                    underline=font_kwargs.get("underline", cell.font.underline),
                    color=font_kwargs.get("color", cell.font.color),
                )
            if fill is not None:
                cell.fill = fill
            if border is not None:
                cell.border = border
            if alignment_kwargs:
                cell.alignment = Alignment(
                    horizontal=alignment_kwargs.get("horizontal", cell.alignment.horizontal),
                    vertical=alignment_kwargs.get("vertical", cell.alignment.vertical),
                    wrap_text=alignment_kwargs.get("wrap_text", cell.alignment.wrap_text),
                )
            if op.get("numberFormat"):
                cell.number_format = op["numberFormat"]


def _clear(workbook, op: dict, formats: bool) -> None:
    ws = _sheet(workbook, op["sheet"])
    min_col, min_row, max_col, max_row = range_boundaries(str(op["range"]))
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.value = None
            if formats:
                cell.font = Font()
                cell.fill = PatternFill()
                cell.border = Border()
                cell.alignment = Alignment()
                cell.number_format = "General"


def _hex(value: str) -> str:
    text = str(value).strip().lstrip("#")
    if len(text) == 6:
        return "FF" + text.upper()
    return text.upper()


def _border_style(name: str) -> str:
    mapping = {
        "continuous": "thin",
        "thin": "thin",
        "medium": "medium",
        "thick": "thick",
        "double": "double",
        "dash": "dashed",
        "dot": "dotted",
        "hairline": "hair",
    }
    return mapping.get(str(name), str(name))


def read_page(path: str, sheet: str | None, address: str | None, include_styles: bool, limit: int) -> dict:
    path = canonical_path(path)
    workbook = load_workbook(path, data_only=False, read_only=False)
    cached = load_workbook(path, data_only=True, read_only=False)
    try:
        sheet_name = sheet or workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise KeyError(f"sheet not found: {sheet_name}")
        ws = workbook[sheet_name]
        cached_ws = cached[sheet_name]
        if address:
            min_col, min_row, max_col, max_row = range_boundaries(address)
        else:
            min_row, min_col = 1, 1
            max_row = max(ws.max_row or 1, 1)
            max_col = max(ws.max_column or 1, 1)
        values = []
        formulas = []
        styles = []
        next_range = None
        consumed = 0
        stop_row = min_row
        stop_col = min_col
        for row in range(min_row, max_row + 1):
            if consumed >= limit:
                next_range = f"{ws.cell(row, min_col).coordinate}:{ws.cell(max_row, max_col).coordinate}"
                break
            value_row = []
            formula_row = []
            style_row = []
            for col in range(min_col, max_col + 1):
                if consumed >= limit:
                    next_range = f"{ws.cell(row, col).coordinate}:{ws.cell(max_row, max_col).coordinate}"
                    break
                cell = ws.cell(row, col)
                cached_cell = cached_ws.cell(row, col)
                raw = cell.value
                formula = raw if isinstance(raw, str) and str(raw).startswith("=") else None
                formula_row.append(formula)
                value_row.append(cached_cell.value if formula else raw)
                if include_styles:
                    style_row.append(_style_of(cell))
                consumed += 1
                stop_row, stop_col = row, col
            values.append(value_row)
            formulas.append(formula_row)
            if include_styles:
                styles.append(style_row)
            if next_range:
                break
        start = ws.cell(min_row, min_col).coordinate
        end = ws.cell(stop_row, stop_col).coordinate
        shown = start if start == end else f"{start}:{end}"
        payload = {
            "sheet": sheet_name,
            "range": shown,
            "values": values,
            "formulas": formulas,
            "nextRange": next_range,
            "sheets": workbook.sheetnames,
        }
        if include_styles:
            payload["styles"] = styles
        return payload
    finally:
        workbook.close()
        cached.close()


def _style_of(cell) -> dict:
    font = cell.font
    fill = cell.fill
    alignment = cell.alignment
    return {
        "fontName": font.name,
        "fontSize": font.size,
        "bold": bool(font.bold),
        "italic": bool(font.italic),
        "underline": bool(font.underline),
        "fontColor": _color_out(font.color),
        "fillColor": _color_out(getattr(fill.fgColor, "rgb", None) and fill.fgColor or fill.fgColor),
        "numberFormat": cell.number_format,
        "horizontalAlignment": alignment.horizontal,
        "verticalAlignment": alignment.vertical,
        "wrapText": bool(alignment.wrap_text),
    }


def _color_out(color) -> str | dict | None:
    if color is None:
        return None
    color_type = getattr(color, "type", None)
    if color_type == "rgb" and color.rgb and color.rgb != "00000000":
        rgb = str(color.rgb)
        return "#" + rgb[-6:]
    if color_type == "theme":
        return {"theme": color.theme, "tint": color.tint}
    if color_type == "indexed":
        return {"indexed": color.indexed}
    return None


def discard(path: str | None) -> None:
    if path and os.path.exists(path):
        try:
            os.remove(path)
        except OSError:
            pass


def file_hash(path: str) -> str:
    import hashlib

    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def copy_file(src: str, dest: str) -> None:
    shutil.copy2(src, dest)
