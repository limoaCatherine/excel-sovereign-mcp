"""Batch routing. The tool name never selects the engine."""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

from excel_sovereign.inspect import PackageInfo, range_size, ranges_overlap

MAX_WRITE_CELLS = 100_000

STYLE_ACTIONS = {
    "format",
    "set_row_height",
    "set_column_width",
    "merge",
    "unmerge",
}
VALUE_ACTIONS = {"set_values", "clear_contents", "clear_all"}
FORMULA_ACTIONS = {"set_formulas", "define_name"}
OPENPYXL_SHEET_ACTIONS = {"create_sheet", "create_workbook"}
STRUCTURE_ACTIONS = {
    "rename_sheet",
    "copy_sheet",
    "move_sheet",
    "delete_sheet",
    "insert_rows",
    "delete_rows",
    "insert_columns",
    "delete_columns",
    "insert_cells",
    "delete_cells",
}
ROW_COL_ACTIONS = {
    "insert_rows",
    "delete_rows",
    "insert_columns",
    "delete_columns",
    "insert_cells",
    "delete_cells",
}

# action -> excelcli command. In-session commands do not save by themselves.
COM_COMMANDS: dict[str, str] = {
    "set_values": "range.set-values",
    "set_formulas": "range.set-formulas",
    "clear_contents": "range.clear-contents",
    "clear_all": "range.clear-all",
    "format": "rangeformat.format-ranges",
    "merge": "rangeformat.merge-cells",
    "unmerge": "rangeformat.unmerge-cells",
    "set_row_height": "rangeformat.set-row-height",
    "set_column_width": "rangeformat.set-column-width",
    "define_name": "namedrange.create",
    "create_sheet": "sheet.create",
    "rename_sheet": "sheet.rename",
    "copy_sheet": "sheet.copy",
    "move_sheet": "sheet.move",
    "delete_sheet": "sheet.delete",
    "insert_rows": "rangeedit.insert-rows",
    "delete_rows": "rangeedit.delete-rows",
    "insert_columns": "rangeedit.insert-columns",
    "delete_columns": "rangeedit.delete-columns",
    "insert_cells": "rangeedit.insert-cells",
    "delete_cells": "rangeedit.delete-cells",
    "table_list": "table.list",
    "table_create": "table.create",
    "table_rename": "table.rename",
    "table_delete": "table.delete",
    "table_resize": "table.resize",
    "table_append": "table.append",
    "table_apply_filter": "tablecolumn.apply-filter",
    "table_clear_filters": "tablecolumn.clear-filters",
    "table_add_to_data_model": "table.add-to-data-model",
    "table_create_from_dax": "table.create-from-dax",
    "table_set_style": "table.set-style",
    "table_read": "table.read",
    "pivot_list": "pivottable.list",
    "pivot_create_from_range": "pivottable.create-from-range",
    "pivot_create_from_table": "pivottable.create-from-table",
    "pivot_delete": "pivottable.delete",
    "pivot_refresh": "pivottable.refresh",
    "chart_list": "chart.list",
    "chart_create_from_range": "chart.create-from-range",
    "chart_create_from_table": "chart.create-from-table",
    "chart_delete": "chart.delete",
    "chart_move": "chart.move",
    "chart_fit": "chart.fit-to-range",
    "powerquery_list": "powerquery.list",
    "powerquery_view": "powerquery.view",
    "powerquery_create": "powerquery.create",
    "powerquery_update": "powerquery.update",
    "powerquery_refresh": "powerquery.refresh",
    "powerquery_refresh_all": "powerquery.refresh-all",
    "powerquery_delete": "powerquery.delete",
    "powerquery_rename": "powerquery.rename",
    "datamodel_list_tables": "datamodel.list-tables",
    "datamodel_list_measures": "datamodel.list-measures",
    "datamodel_create_measure": "datamodel.create-measure",
    "datamodel_update_measure": "datamodel.update-measure",
    "datamodel_delete_measure": "datamodel.delete-measure",
    "datamodel_evaluate": "datamodel.evaluate",
    "datamodel_refresh": "datamodel.refresh",
    "conditional_format_add": "conditionalformat.add-rule",
    "conditional_format_clear": "conditionalformat.clear-rules",
    "conditional_format_list": "conditionalformat.list-rules",
    "validation_add": "rangeformat.validate-range",
    "validation_get": "rangeformat.get-validation",
    "validation_remove": "rangeformat.remove-validation",
    "comment_set": "sheet.set-comment",
    "comment_get": "sheet.get-comment",
    "comment_clear": "sheet.clear-comment",
    "threaded_comment_add": "rangelink.add-threaded-comment",
    "hyperlink_add": "rangelink.add-hyperlink",
    "hyperlink_remove": "rangelink.remove-hyperlink",
    "freeze": "window.freeze-panes",
    "unfreeze": "window.unfreeze-panes",
    "sheet_hide": "sheet.hide",
    "sheet_show": "sheet.show",
    "vba_list": "vba.list",
    "vba_view": "vba.view",
    "vba_import": "vba.import",
    "vba_update": "vba.update",
    "vba_run": "vba.run",
    "vba_delete": "vba.delete",
    "slicer_create": "slicer.create-slicer",
    "slicer_list": "slicer.list-slicers",
    "slicer_delete": "slicer.delete-slicer",
}

VBA_ACTIONS = {name for name in COM_COMMANDS if name.startswith("vba_")}
COM_COMMANDS["define_name_update"] = "namedrange.update"
# Expanded into delete_rows / delete_columns before routing.
EXPANDED_ACTIONS = {"trim_sheet"}
KNOWN_ACTIONS = set(COM_COMMANDS) | {"verify", "create_workbook"} | EXPANDED_ACTIONS

# These return data and change nothing. A call made only of them opens, reads, and closes without saving.
QUERY_ACTIONS = {
    "table_list",
    "table_read",
    "pivot_list",
    "chart_list",
    "slicer_list",
    "powerquery_list",
    "powerquery_view",
    "datamodel_list_tables",
    "datamodel_list_measures",
    "datamodel_evaluate",
    "conditional_format_list",
    "validation_get",
    "comment_get",
    "vba_list",
    "vba_view",
}

_MODEL_TABLE = {"table_add_to_data_model", "table_create_from_dax"}
_TABLE_PREFIXES = ("table_", "pivot_", "chart_", "slicer_")
_MODEL_PREFIXES = ("powerquery_", "datamodel_")
_VIEW_ACTIONS = {
    "conditional_format_add",
    "conditional_format_clear",
    "conditional_format_list",
    "validation_add",
    "validation_get",
    "validation_remove",
    "comment_set",
    "comment_get",
    "comment_clear",
    "threaded_comment_add",
    "hyperlink_add",
    "hyperlink_remove",
    "freeze",
    "unfreeze",
    "sheet_hide",
    "sheet_show",
}

# Cell, style, and sheet-structure actions ride along on every tool, so seeding data and
# building a table on it is one Excel session instead of two.
CORE_ACTIONS = (
    STYLE_ACTIONS
    | VALUE_ACTIONS
    | FORMULA_ACTIONS
    | OPENPYXL_SHEET_ACTIONS
    | STRUCTURE_ACTIONS
    | {"define_name_update", "verify"}
)

# Tool name selects the action family. It does not select the engine.
TOOL_ACTIONS: dict[str, set[str]] = {
    "workbook_apply": CORE_ACTIONS | EXPANDED_ACTIONS | {"layout"},
    "excel_table": {
        name
        for name in COM_COMMANDS
        if name.startswith(_TABLE_PREFIXES) and name not in _MODEL_TABLE
    }
    | CORE_ACTIONS,
    "excel_model": {name for name in COM_COMMANDS if name.startswith(_MODEL_PREFIXES)}
    | _MODEL_TABLE
    | CORE_ACTIONS,
    "excel_view": set(_VIEW_ACTIONS) | CORE_ACTIONS,
    "excel_vba": set(VBA_ACTIONS) | CORE_ACTIONS,
}


def _owners() -> dict[str, str]:
    owners: dict[str, str] = {}
    for name, actions in TOOL_ACTIONS.items():
        for action in actions:
            owners.setdefault(action, name)
    return owners


def action_outside(tool: str, ops: list[dict]) -> tuple[str, str] | None:
    """Return an action that this tool does not own, and the tool that does."""
    allowed = TOOL_ACTIONS.get(tool)
    if allowed is None:
        return None
    owners = _owners()
    for op in ops:
        action = op["action"]
        if action in allowed or action not in owners:
            continue
        return action, owners[action]
    return None


def is_query(ops: list[dict]) -> bool:
    return bool(ops) and all(op["action"] in QUERY_ACTIONS for op in ops)

# These excelcli actions open files and save outside the session. They are not mapped.
IMMEDIATE_SAVE_BLOCKED = {"sheet.copy-to-file", "sheet.move-to-file"}


@dataclass
class RouteDecision:
    engine: str
    calc_required: bool
    allow_macros: bool
    reason: str


@dataclass
class WriteSpan:
    sheet: str
    range: str
    formulas: list[str] = field(default_factory=list)


def normalize_ops(ops: list[dict]) -> list[dict]:
    if not isinstance(ops, list) or not ops:
        raise ValueError("ops must be a non-empty list")
    normalized = []
    for index, op in enumerate(ops):
        if not isinstance(op, dict) or "action" not in op:
            raise ValueError(f"ops[{index}] needs an action")
        item = dict(op)
        item["action"] = str(item["action"]).replace("-", "_")
        normalized.append(item)
    return normalized


def validate_known(ops: list[dict]) -> str | None:
    for op in ops:
        action = op["action"]
        if action not in KNOWN_ACTIONS:
            return action
    return None


def count_cells(ops: list[dict]) -> int:
    total = 0
    for op in ops:
        action = op["action"]
        if action == "set_values":
            values = op.get("values") or []
            total += sum(len(row) for row in values if isinstance(row, list))
        elif action == "set_formulas":
            formulas = op.get("formulas") or []
            total += sum(len(row) for row in formulas if isinstance(row, list))
        elif op.get("range"):
            try:
                total += range_size(str(op["range"]))[2]
            except ValueError:
                continue
    return total


def _wrote_formula(ops: list[dict]) -> bool:
    if any(op["action"] in FORMULA_ACTIONS or op["action"] == "define_name_update" for op in ops):
        return True
    for op in ops:
        if op["action"] != "set_values":
            continue
        for row in op.get("values") or []:
            if not isinstance(row, list):
                continue
            for value in row:
                if isinstance(value, str) and value.startswith("="):
                    return True
    return False


def _structural(ops: list[dict]) -> bool:
    return any(op["action"] in STRUCTURE_ACTIONS for op in ops)


def _value_change(ops: list[dict]) -> bool:
    return any(op["action"] in VALUE_ACTIONS or op["action"] in FORMULA_ACTIONS or op["action"] == "define_name_update" for op in ops)


def _style_only(ops: list[dict]) -> bool:
    return all(op["action"] in STYLE_ACTIONS for op in ops)


def _table_conflict(ops: list[dict], info: PackageInfo) -> bool:
    for op in ops:
        sheet = str(op.get("sheet") or "")
        for table in info.tables:
            same_sheet = not table.sheet or table.sheet == sheet
            if not same_sheet:
                continue
            if op["action"] in ROW_COL_ACTIONS:
                return True
            target = op.get("range")
            if target and table.ref and ranges_overlap(str(target), table.ref):
                return True
    return False


def decide(ops: list[dict], info: PackageInfo | None) -> RouteDecision:
    info = info or PackageInfo()
    if any(op["action"] == "verify" for op in ops):
        return RouteDecision("com", True, False, "verify")
    allow_macros = any(op["action"] in VBA_ACTIONS for op in ops)
    if is_query(ops):
        return RouteDecision("com", False, allow_macros, "query")
    openpyxl_actions = STYLE_ACTIONS | VALUE_ACTIONS | FORMULA_ACTIONS | OPENPYXL_SHEET_ACTIONS
    needs_com = (
        info.has_com_parts
        or _structural(ops)
        or _table_conflict(ops, info)
        or any(op["action"] not in openpyxl_actions for op in ops)
    )
    wrote_formula = _wrote_formula(ops)
    has_formula_book = info.has_calc_chain or info.has_formulas or info.has_defined_names or wrote_formula
    changes_values = _value_change(ops) or _structural(ops) or any(op["action"] == "create_sheet" for op in ops)
    if _style_only(ops):
        calc = False
    elif changes_values:
        calc = has_formula_book or wrote_formula
    else:
        calc = False
    engine = "com" if needs_com else "openpyxl"
    reason = "com-parts" if info.has_com_parts and engine == "com" else engine
    return RouteDecision(engine, calc, allow_macros, reason)


def shift_address(address: str, op: dict) -> str:
    """Move an address that was recorded before a structural op on the same sheet."""
    action = op["action"]
    sheet = str(op.get("sheet") or "")
    if ":" in address:
        min_col, min_row, max_col, max_row = range_boundaries(address)
    else:
        letters, row = coordinate_from_string(address.replace("$", ""))
        min_col = max_col = column_index_from_string(letters)
        min_row = max_row = row
    if action == "insert_rows":
        row = int(op["row"])
        count = int(op.get("count") or 1)
        if min_row >= row:
            min_row += count
            max_row += count
    elif action == "delete_rows":
        row = int(op["row"])
        count = int(op.get("count") or 1)
        if min_row >= row + count:
            min_row -= count
            max_row -= count
    elif action == "insert_columns":
        col = column_index_from_string(str(op["column"]))
        count = int(op.get("count") or 1)
        if min_col >= col:
            min_col += count
            max_col += count
    elif action == "delete_columns":
        col = column_index_from_string(str(op["column"]))
        count = int(op.get("count") or 1)
        if min_col >= col + count:
            min_col -= count
            max_col -= count
    del sheet
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def matrix_range(address: str, row_count: int, col_count: int) -> str:
    """Size a COM write to the matrix. excelcli rejects a range whose shape differs."""
    min_col, min_row, _, _ = range_boundaries(address)
    if row_count < 1 or col_count < 1:
        return address
    start = f"{get_column_letter(min_col)}{min_row}"
    if row_count == 1 and col_count == 1:
        return start
    end_col = get_column_letter(min_col + col_count - 1)
    end_row = min_row + row_count - 1
    return f"{start}:{end_col}{end_row}"


def com_args(op: dict) -> dict:
    """Translate one op into excelcli arguments. Does not include the session."""
    action = op["action"]
    extra = dict(op.get("args") or {})
    sheet = op.get("sheet")
    rng = op.get("range")
    if action == "set_values":
        values = op.get("values") or []
        rows = len(values) if isinstance(values, list) else 0
        cols = max((len(row) for row in values if isinstance(row, list)), default=0)
        address = matrix_range(str(rng), rows, cols) if rng and rows and cols else rng
        return {"sheetName": sheet, "rangeAddress": address, "values": values, **extra}
    if action == "set_formulas":
        formulas = op.get("formulas") or []
        rows = len(formulas) if isinstance(formulas, list) else 0
        cols = max((len(row) for row in formulas if isinstance(row, list)), default=0)
        address = matrix_range(str(rng), rows, cols) if rng and rows and cols else rng
        return {"sheetName": sheet, "rangeAddress": address, "formulas": formulas, **extra}
    if action in {"clear_contents", "clear_all", "merge", "unmerge"}:
        return {"sheetName": sheet, "rangeAddress": rng, **extra}
    if action == "format":
        payload = {
            "sheetName": sheet,
            "rangeAddresses": [rng],
            "fontName": op.get("fontName"),
            "fontSize": op.get("fontSize"),
            "bold": op.get("bold"),
            "italic": op.get("italic"),
            "underline": op.get("underline"),
            "fontColor": op.get("fontColor"),
            "fillColor": op.get("fillColor"),
            "borderStyle": op.get("borderStyle"),
            "borderColor": op.get("borderColor"),
            "borderWeight": op.get("borderWeight"),
            "horizontalAlignment": op.get("horizontalAlignment"),
            "verticalAlignment": op.get("verticalAlignment"),
            "wrapText": op.get("wrapText"),
            "numberFormat": op.get("numberFormat"),
        }
        payload = {key: value for key, value in payload.items() if value is not None}
        payload.update(extra)
        return payload
    if action == "set_row_height":
        target = rng or f"{op['row']}:{int(op['row']) + int(op.get('count') or 1) - 1}"
        return {"sheetName": sheet, "rangeAddress": target, "rowHeight": op.get("height"), **extra}
    if action == "set_column_width":
        target = rng or f"{op['column']}:{op['column']}"
        return {"sheetName": sheet, "rangeAddress": target, "columnWidth": op.get("width"), **extra}
    if action in {"define_name", "define_name_update"}:
        reference = op.get("formula") or op.get("reference")
        return {"name": op.get("name"), "reference": reference, **extra}
    if action == "create_sheet":
        return {"sheetName": op.get("name") or sheet, **extra}
    if action == "rename_sheet":
        return {"oldName": op.get("oldName") or sheet, "newName": op.get("newName"), **extra}
    if action == "copy_sheet":
        return {"sourceName": op.get("source") or sheet, "targetName": op.get("newName"), **extra}
    if action == "move_sheet":
        return {
            "sheetName": sheet,
            "beforeSheet": op.get("before"),
            "afterSheet": op.get("after"),
            **extra,
        }
    if action == "delete_sheet":
        return {"sheetName": sheet or op.get("name"), **extra}
    if action in {"insert_rows", "delete_rows"}:
        row = int(op["row"])
        count = int(op.get("count") or 1)
        address = rng or f"{row}:{row + count - 1}"
        return {"sheetName": sheet, "rangeAddress": address, **extra}
    if action in {"insert_columns", "delete_columns"}:
        column = str(op["column"])
        count = int(op.get("count") or 1)
        end = get_column_letter(column_index_from_string(column) + count - 1)
        address = rng or f"{column}:{end}"
        return {"sheetName": sheet, "rangeAddress": address, **extra}
    if action == "insert_cells":
        return {"sheetName": sheet, "rangeAddress": rng, "insertShift": op.get("shift") or "Down", **extra}
    if action == "delete_cells":
        return {"sheetName": sheet, "rangeAddress": rng, "deleteShift": op.get("shift") or "Up", **extra}
    if action == "comment_set":
        return {"sheetName": sheet, "cellAddress": op.get("cell") or rng, "text": op.get("text"), **extra}
    if action in {"comment_get", "comment_clear"}:
        return {"sheetName": sheet, "cellAddress": op.get("cell") or rng, **extra}
    if action == "freeze":
        return {
            "sheetName": sheet,
            "frozenRows": op.get("rows") or 0,
            "frozenColumns": op.get("columns") or 0,
            **extra,
        }
    if action == "unfreeze":
        return {"sheetName": sheet, **extra}
    if action in {"sheet_hide", "sheet_show"}:
        return {"sheetName": sheet, **extra}
    if action == "vba_run":
        return {"procedureName": op.get("procedure"), "parameters": op.get("parameters") or [], **extra}
    if action in {"vba_view", "vba_delete"}:
        return {"moduleName": op.get("module"), **extra}
    if action in {"vba_import", "vba_update"}:
        return {"moduleName": op.get("module"), "vbaCode": op.get("code"), **extra}
    if extra:
        return extra
    translated = {}
    for key, value in op.items():
        if key in {"action", "args"}:
            continue
        translated[key] = value
    if "sheet" in translated:
        translated["sheetName"] = translated.pop("sheet")
    if "range" in translated:
        translated["rangeAddress"] = translated.pop("range")
    return translated
