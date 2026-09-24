"""Sheet layout. Domains differ by which blocks they stack, not by palette.

A profile name does not select colors and does not forbid blocks. Finance, analytics,
games, and anything else share the same tokens. The caller supplies the blocks that
match the actual columns.
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter

TOKENS = {
    "font": "微软雅黑",
    "title": {"size": 16, "color": "#1F4E79", "fill": "#F2F2F2", "height": 26},
    "section": {"size": 12, "color": "#1F4E79", "fill": "#E7E6E6", "height": 22},
    "subsection": {"size": 10, "color": "#404040", "fill": "#F2F2F2", "height": 18},
    "header": {"size": 10, "color": "#404040", "fill": "#D9D9D9", "height": 18},
    "data": {"size": 10, "height": 18, "fill": "#FFFFFF"},
    "input_fill": "#F3F8FC",
    "input_border": "#D8E6F0",
    "input_font": "#0000FF",
    "cross_sheet": "#008000",
    "structured": "#FF0000",
    "formula": "#000000",
    "label": "#000000",
    "empty": "#808080",
    "group_line": "#000000",
}

WIDTHS = {
    "spacer": 3,
    "number": 12,
    "percent": 12,
    "id": 10,
    "date": 12,
    "unit": 8,
    "label": 18,
    "enum": 14,
    "text": 28,
    "formula": 28,
    "group": 14,
}

FORMATS = {
    "id": "0",
    "number": "0.000000;-0.000000;0.000000",
    "percent": "0.00%",
    "date": "yyyy-mm-dd",
}

ROLES = set(WIDTHS)
CONTENT_KINDS = {"parameters", "catalog", "formula", "fact"}

# Recommended assemblies. Any profile may still contain any block.
PROFILES = {
    "finance": ("title", "section", "parameters", "formula", "catalog"),
    "analytics": ("title", "section", "parameters", "fact", "catalog"),
    "general": ("title", "section", "subsection", "parameters", "catalog", "formula", "fact"),
}


def expand_layouts(ops: list[dict]) -> list[dict]:
    expanded = []
    for op in ops:
        if op.get("action") == "layout":
            expanded.extend(compile_layout(op))
        else:
            expanded.append(op)
    if not expanded:
        raise ValueError("ops must be a non-empty list")
    return expanded


def compile_layout(op: dict) -> list[dict]:
    sheet = str(op.get("sheet") or "")
    if not sheet:
        raise ValueError("layout needs a sheet")
    profile = op.get("profile")
    if profile is not None and profile not in PROFILES:
        raise ValueError(f"unknown layout profile: {profile}")
    blocks = op.get("blocks")
    if not isinstance(blocks, list) or not blocks:
        raise ValueError("layout needs blocks")
    sheet_width = max(_content_width(block) for block in blocks)
    sheet_width = max(sheet_width, 1)
    ops: list[dict] = []
    if op.get("create"):
        ops.append({"action": "create_sheet", "name": sheet})
    row = 1
    widths: dict[str, float] = {}
    for index, block in enumerate(blocks):
        if index:
            row += 1
        emitted, row = _compile_block(sheet, block, row, sheet_width, widths)
        ops.extend(emitted)
    for column, width in widths.items():
        ops.append({"action": "set_column_width", "sheet": sheet, "column": column, "width": width})
    return ops


def _content_width(block: dict) -> int:
    kind = str(block.get("kind") or "")
    if kind == "formula":
        bands = block.get("bands") or []
        total = sum(len(band.get("columns") or []) for band in bands)
        if len(bands) > 1:
            total += len(bands) - 1
        return total
    if kind in CONTENT_KINDS:
        return len(block.get("columns") or [])
    return int(block.get("span") or 0)


def _compile_block(sheet: str, block: dict, row: int, sheet_width: int, widths: dict[str, float]):
    kind = str(block.get("kind") or "")
    if kind == "title":
        return _heading(sheet, block, row, sheet_width, "title", widths)
    if kind == "section":
        return _heading(sheet, block, row, int(block.get("span") or sheet_width), "section", widths)
    if kind == "subsection":
        return _heading(sheet, block, row, 1, "subsection", widths)
    if kind == "parameters":
        return _grid(sheet, block, row, widths, grouped=False)
    if kind == "catalog":
        return _grid(sheet, block, row, widths, grouped=True)
    if kind == "formula":
        return _formula(sheet, block, row, widths)
    if kind == "fact":
        return _fact(sheet, block, row, widths)
    raise ValueError(f"unknown layout block: {kind}")


def _heading(sheet: str, block: dict, row: int, span: int, level: str, widths: dict[str, float]):
    span = max(span, 1)
    token = TOKENS[level]
    address = _range(1, row, span, row)
    ops = [
        {"action": "set_values", "sheet": sheet, "range": _cell(1, row), "values": [[block.get("text") or ""]]},
        _paint(
            sheet,
            address,
            size=token["size"],
            color=token["color"],
            fill=token["fill"],
            bold=True,
        ),
        {"action": "set_row_height", "sheet": sheet, "row": row, "height": token["height"]},
    ]
    _remember_width(widths, 1, "label")
    return ops, row + 1


def _grid(sheet: str, block: dict, row: int, widths: dict[str, float], grouped: bool):
    columns = _columns(block)
    rows = _rows(block, len(columns))
    header_row = row
    data_start = row + 1
    ops = [_values(sheet, 1, header_row, [[column["header"] for column in columns], *rows])]
    ops.append(_paint(sheet, _range(1, header_row, len(columns), header_row), header=True))
    ops.append({"action": "set_row_height", "sheet": sheet, "row": header_row, "height": TOKENS["header"]["height"]})
    for index, column in enumerate(columns):
        _remember_width(widths, index + 1, column["role"], column.get("width"))
        if not rows:
            continue
        number_format = FORMATS.get(column["role"])
        if number_format:
            ops.append(
                _paint(
                    sheet,
                    _range(index + 1, data_start, index + 1, data_start + len(rows) - 1),
                    number_format=number_format,
                )
            )
    group_index = _group_index(block, columns) if grouped else None
    previous = None
    for r_index, values in enumerate(rows):
        data_row = data_start + r_index
        ops.append({"action": "set_row_height", "sheet": sheet, "row": data_row, "height": TOKENS["data"]["height"]})
        for c_index, column in enumerate(columns):
            value = values[c_index]
            ops.append(
                _paint(
                    sheet,
                    _cell(c_index + 1, data_row),
                    **_cell_style(column, value),
                )
            )
        if group_index is not None:
            if previous is not None and values[group_index] != previous:
                ops.append(
                    _paint(
                        sheet,
                        _range(1, data_row, len(columns), data_row),
                        border_style="thin",
                        border_color=TOKENS["group_line"],
                        border_edges="top",
                    )
                )
            previous = values[group_index]
    return ops, data_start + len(rows)


def _formula(sheet: str, block: dict, row: int, widths: dict[str, float]):
    bands = block.get("bands")
    if not isinstance(bands, list) or not bands:
        raise ValueError("formula block needs bands")
    ops = []
    col = 1
    data_rows = 0
    for index, band in enumerate(bands):
        if index:
            _remember_width(widths, col, "spacer")
            col += 1
        columns = _columns(band)
        rows = _rows(band, len(columns))
        data_rows = max(data_rows, len(rows))
        start = col
        end = col + len(columns) - 1
        ops.append(_values(sheet, start, row, [[band.get("title") or ""]]))
        ops.append(_paint(sheet, _range(start, row, end, row), level="section"))
        ops.append(_values(sheet, start, row + 1, [[column["header"] for column in columns], *rows]))
        ops.append(_paint(sheet, _range(start, row + 1, end, row + 1), header=True))
        for c_index, column in enumerate(columns):
            _remember_width(widths, start + c_index, column["role"], column.get("width"))
            number_format = FORMATS.get(column["role"])
            if number_format and rows:
                ops.append(
                    _paint(
                        sheet,
                        _range(start + c_index, row + 2, start + c_index, row + 1 + len(rows)),
                        number_format=number_format,
                    )
                )
            for r_index, values in enumerate(rows):
                ops.append(
                    _paint(
                        sheet,
                        _cell(start + c_index, row + 2 + r_index),
                        **_cell_style(column, values[c_index]),
                    )
                )
        col = end + 1
    ops.append({"action": "set_row_height", "sheet": sheet, "row": row, "height": TOKENS["section"]["height"]})
    ops.append({"action": "set_row_height", "sheet": sheet, "row": row + 1, "height": TOKENS["header"]["height"]})
    for offset in range(data_rows):
        ops.append(
            {"action": "set_row_height", "sheet": sheet, "row": row + 2 + offset, "height": TOKENS["data"]["height"]}
        )
    return ops, row + 2 + data_rows


def _fact(sheet: str, block: dict, row: int, widths: dict[str, float]):
    name = str(block.get("name") or "")
    if not name:
        raise ValueError("fact block needs a table name")
    columns = _columns(block)
    rows = _rows(block, len(columns))
    last = row + len(rows)
    ops = [_values(sheet, 1, row, [[column["header"] for column in columns], *rows])]
    ops.append(
        {
            "action": "table_create",
            "args": {
                "sheetName": sheet,
                "tableName": name,
                "rangeAddress": _range(1, row, len(columns), last),
                "hasHeaders": True,
            },
        }
    )
    for index, column in enumerate(columns):
        _remember_width(widths, index + 1, column["role"], column.get("width"))
    return ops, last + 1


def _columns(block: dict) -> list[dict]:
    columns = block.get("columns")
    if not isinstance(columns, list) or not columns:
        raise ValueError(f"{block.get('kind')} block needs columns")
    normalized = []
    for column in columns:
        if not isinstance(column, dict) or "header" not in column:
            raise ValueError("each column needs a header")
        role = str(column.get("role") or "text")
        if role not in ROLES:
            raise ValueError(f"unknown column role: {role}")
        item = dict(column)
        item["role"] = role
        item["header"] = "" if role == "spacer" else column.get("header")
        normalized.append(item)
    return normalized


def _rows(block: dict, width: int) -> list[list]:
    rows = block.get("rows") or []
    if not isinstance(rows, list):
        raise ValueError("rows must be a list")
    padded = []
    for row in rows:
        if not isinstance(row, list):
            raise ValueError("each row must be a list")
        values = list(row[:width])
        while len(values) < width:
            values.append(None)
        padded.append(values)
    return padded


def _group_index(block: dict, columns: list[dict]) -> int | None:
    name = block.get("group")
    if not name:
        for index, column in enumerate(columns):
            if column["role"] == "group":
                return index
        return None
    for index, column in enumerate(columns):
        if column["header"] == name:
            return index
    raise ValueError(f"catalog group column not found: {name}")


def _cell_style(column: dict, value) -> dict:
    color = _font_color(column, value)
    style = {
        "size": TOKENS["data"]["size"],
        "color": color,
        "fill": TOKENS["data"]["fill"],
        "bold": False,
    }
    if column["role"] == "text":
        style["wrap"] = True
    if _is_input(column, value):
        style["fill"] = TOKENS["input_fill"]
        style["border_style"] = "thin"
        style["border_color"] = TOKENS["input_border"]
    return style


def _font_color(column: dict, value) -> str:
    if value is None or value == "" or value == "—":
        return TOKENS["empty"]
    if isinstance(value, str) and value.startswith("="):
        if "[" in value:
            return TOKENS["structured"]
        if "!" in value:
            return TOKENS["cross_sheet"]
        return TOKENS["formula"]
    if _is_input(column, value):
        return TOKENS["input_font"]
    return TOKENS["label"]


def _is_input(column: dict, value) -> bool:
    if not column.get("input"):
        return False
    if value is None or value == "" or value == "—":
        return False
    return not (isinstance(value, str) and value.startswith("="))


def _values(sheet: str, col: int, row: int, values: list[list]) -> dict:
    return {"action": "set_values", "sheet": sheet, "range": _cell(col, row), "values": values}


def _paint(sheet: str, address: str, **style) -> dict:
    token_size = style.get("size", TOKENS["data"]["size"])
    payload = {
        "action": "format",
        "sheet": sheet,
        "range": address,
        "fontName": TOKENS["font"],
        "fontSize": token_size,
        "horizontalAlignment": "left",
        "verticalAlignment": "center",
    }
    if "color" in style:
        payload["fontColor"] = style["color"]
    if "fill" in style:
        payload["fillColor"] = style["fill"]
    if "bold" in style:
        payload["bold"] = style["bold"]
    if style.get("header"):
        header = TOKENS["header"]
        payload["fontSize"] = header["size"]
        payload["fontColor"] = header["color"]
        payload["fillColor"] = header["fill"]
        payload["bold"] = True
    if style.get("level"):
        token = TOKENS[style["level"]]
        payload["fontSize"] = token["size"]
        payload["fontColor"] = token["color"]
        payload["fillColor"] = token["fill"]
        payload["bold"] = True
    if style.get("wrap"):
        payload["wrapText"] = True
    if style.get("number_format"):
        payload["numberFormat"] = style["number_format"]
    if style.get("border_style"):
        payload["borderStyle"] = style["border_style"]
        payload["borderColor"] = style.get("border_color") or "#000000"
        if style.get("border_edges"):
            payload["borderEdges"] = style["border_edges"]
    return payload


def _remember_width(widths: dict[str, float], col: int, role: str, explicit=None) -> None:
    letter = get_column_letter(col)
    width = float(explicit if explicit is not None else WIDTHS[role])
    widths[letter] = max(widths.get(letter, 0), width)


def _cell(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"


def _range(c1: int, r1: int, c2: int, r2: int) -> str:
    start = _cell(c1, r1)
    end = _cell(c2, r2)
    return start if start == end else f"{start}:{end}"
