"""Zip inspection: COM-owned parts, tables, names, and cross-sheet formula refs."""

from __future__ import annotations

import posixpath
import re
import zipfile
from dataclasses import dataclass, field
from xml.etree import ElementTree

from openpyxl.utils import range_boundaries
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REL_NS = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
PKG_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

COM_PREFIXES = (
    "xl/pivotCaches/",
    "xl/pivotTables/",
    "xl/queryTables/",
    "xl/model/",
    "customXml/",
    "xl/slicers/",
    "xl/slicerCaches/",
    "xl/charts/",
    "xl/richData/",
    "xl/ctrlProps/",
    "xl/activeX/",
    "xl/embeddings/",
)
COM_FILES = {"xl/connections.xml", "xl/metadata.xml"}
OLE2 = bytes([0xD0, 0xCF, 0x11, 0xE0, 0xA1, 0xB1, 0x1A, 0xE1])


@dataclass
class TableRef:
    name: str
    sheet: str
    ref: str


@dataclass
class PackageInfo:
    names: list[str] = field(default_factory=list)
    has_calc_chain: bool = False
    has_formulas: bool = False
    has_defined_names: bool = False
    has_vba: bool = False
    com_parts: list[str] = field(default_factory=list)
    tables: list[TableRef] = field(default_factory=list)

    @property
    def has_com_parts(self) -> bool:
        return bool(self.com_parts)


def is_encrypted_or_irm(path: str) -> bool:
    try:
        with open(path, "rb") as handle:
            signature = handle.read(8)
    except OSError:
        return False
    return signature == OLE2


def inspect_workbook(path: str) -> PackageInfo:
    info = PackageInfo()
    with zipfile.ZipFile(path) as archive:
        names = [name.replace("\\", "/") for name in archive.namelist()]
        info.names = names
        info.has_calc_chain = "xl/calcChain.xml" in names
        info.has_formulas = _has_formulas(archive, names)
        info.has_vba = "xl/vbaProject.bin" in names
        for name in names:
            if name in COM_FILES or any(name.startswith(prefix) for prefix in COM_PREFIXES):
                info.com_parts.append(name)
        info.has_defined_names = _has_defined_names(archive)
        info.tables = _tables(archive, names)
    return info


def _has_formulas(archive: zipfile.ZipFile, names: list[str]) -> bool:
    for name in names:
        if not (name.startswith("xl/worksheets/sheet") and name.endswith(".xml")):
            continue
        if b"<f" in archive.read(name):
            return True
    return False


def _has_defined_names(archive: zipfile.ZipFile) -> bool:
    if "xl/workbook.xml" not in archive.namelist():
        return False
    root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    node = root.find("m:definedNames", NS)
    return node is not None and len(list(node)) > 0


def _sheet_targets(archive: zipfile.ZipFile) -> dict[str, str]:
    """Map worksheet part name to sheet name."""
    root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_by_id = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels}
    sheets: dict[str, str] = {}
    for sheet in root.findall("m:sheets/m:sheet", NS):
        rel_id = sheet.attrib.get(f"{{{PKG_REL}}}id")
        target = rel_by_id.get(rel_id or "", "")
        part = posixpath.normpath(posixpath.join("xl", target)).replace("\\", "/")
        sheets[part] = sheet.attrib.get("name", "")
    return sheets


def _tables(archive: zipfile.ZipFile, names: list[str]) -> list[TableRef]:
    sheet_of_part = _sheet_targets(archive)
    table_sheet: dict[str, str] = {}
    for part, sheet_name in sheet_of_part.items():
        rels_name = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
        rels_name = rels_name.replace("\\", "/")
        if rels_name not in names:
            continue
        rels = ElementTree.fromstring(archive.read(rels_name))
        for rel in rels:
            target = rel.attrib.get("Target", "")
            table_part = posixpath.normpath(posixpath.join(posixpath.dirname(part), target)).replace("\\", "/")
            table_sheet[table_part] = sheet_name
    found: list[TableRef] = []
    for name in names:
        if not name.startswith("xl/tables/table") or not name.endswith(".xml"):
            continue
        root = ElementTree.fromstring(archive.read(name))
        ref = root.attrib.get("ref", "")
        display = root.attrib.get("displayName") or root.attrib.get("name") or name
        found.append(TableRef(name=display, sheet=table_sheet.get(name, ""), ref=ref))
    return found


def ranges_overlap(a: str, b: str) -> bool:
    a1, a2, a3, a4 = range_boundaries(a)
    b1, b2, b3, b4 = range_boundaries(b)
    return not (a3 < b1 or b3 < a1 or a4 < b2 or b4 < a2)


def formula_mentions_sheet(formula: str, sheet: str) -> bool:
    if not formula or not sheet:
        return False
    escaped = re.escape(sheet)
    pattern = rf"(?:'{escaped}'|{escaped})!"
    return re.search(pattern, formula, flags=re.IGNORECASE) is not None


def dependent_ranges(path: str, sheet_names: list[str]) -> dict[str, list[str]]:
    """Cells on other sheets whose formula text mentions one of the given sheets."""
    hits: dict[str, list[str]] = {}
    with zipfile.ZipFile(path) as archive:
        sheets = _sheet_targets(archive)
        for part, sheet in sheets.items():
            if sheet in sheet_names or part not in archive.namelist():
                continue
            try:
                root = ElementTree.fromstring(archive.read(part))
            except ElementTree.ParseError:
                continue
            cells: list[str] = []
            for cell in root.findall(".//m:c", NS):
                formula = cell.find("m:f", NS)
                if formula is None or not formula.text:
                    continue
                if any(formula_mentions_sheet(formula.text, name) for name in sheet_names):
                    ref = cell.attrib.get("r")
                    if ref:
                        cells.append(ref)
            if cells:
                hits[sheet] = cells
    return hits


def bounding_box(cells: list[str]) -> str | None:
    if not cells:
        return None
    min_col = min_row = 10**9
    max_col = max_row = 0
    for cell in cells:
        if ":" in cell:
            c1, r1, c2, r2 = range_boundaries(cell)
        else:
            letters, row = coordinate_from_string(cell.replace("$", ""))
            c1 = c2 = column_index_from_string(letters)
            r1 = r2 = row
        min_col = min(min_col, c1)
        min_row = min(min_row, r1)
        max_col = max(max_col, c2)
        max_row = max(max_row, r2)
    from openpyxl.utils import get_column_letter

    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def range_size(address: str) -> tuple[int, int, int]:
    min_col, min_row, max_col, max_row = range_boundaries(address)
    rows = max_row - min_row + 1
    cols = max_col - min_col + 1
    return rows, cols, rows * cols


def top_left_window(address: str, rows: int = 40, cols: int = 16) -> str:
    from openpyxl.utils import get_column_letter

    min_col, min_row, max_col, max_row = range_boundaries(address)
    end_col = min(max_col, min_col + cols - 1)
    end_row = min(max_row, min_row + rows - 1)
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(end_col)}{end_row}"
