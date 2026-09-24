"""Three tools and the write pipeline. Models do not choose an engine or hold a session."""

from __future__ import annotations

import os
from typing import Any

from mcp.server.mcpserver import Image, MCPServer

from excel_sovereign.book import (
    apply_ops,
    create_empty,
    discard,
    load_for_edit,
    read_page,
    replace_target,
    save_atomic,
)
from excel_sovereign.excel_cli import COM_QUEUE, ComSession, ExcelCliError
from excel_sovereign.inspect import is_encrypted_or_irm, inspect_workbook
from excel_sovereign.layout import expand_layouts
from excel_sovereign.lock import FileLock, LockTimeout, canonical_path, file_is_locked
from excel_sovereign.route import (
    COM_COMMANDS,
    MAX_WRITE_CELLS,
    RouteDecision,
    com_args,
    count_cells,
    decide,
    normalize_ops,
    shift_address,
    validate_known,
)
from excel_sovereign.verify import (
    VerifyOutcome,
    _object_anchors,
    calculate,
    capture_range,
    plan_sheets,
    spill_address,
    union_ranges,
)

mcp = MCPServer("excel-sovereign-mcp")

READ_LIMIT = 4000
ALLOWED = {".xlsx", ".xlsm"}


def _error(code: str, message: str, **extra) -> dict:
    body = {
        "ok": False,
        "committed": False,
        "saved": False,
        "verified": False,
        "calculated": False,
        "calc": "skipped",
        "engine": None,
        "completedOps": 0,
        "failedAt": None,
        "wrote": [],
        "screenshots": [],
        "omitted": [],
        "suggestedNextActions": [],
        "error": {"code": code, "message": message},
    }
    body.update(extra)
    return body


def _extension(path: str) -> str | None:
    suffix = os.path.splitext(path)[1].lower()
    if suffix not in ALLOWED:
        return suffix or "(none)"
    return None


def _response(
    *,
    path: str,
    engine: str,
    tool: str,
    committed: bool,
    calc: str,
    calc_required: bool,
    wrote: list[dict],
    outcome: VerifyOutcome | None,
    completed: int,
    failed_at: int | None,
    error: dict | None,
) -> tuple[dict, list]:
    shots = outcome.shots if outcome else []
    omitted = outcome.omitted if outcome else []
    spill_unknown = bool(outcome and outcome.spill_unknown)
    cropped = any(shot.cropped for shot in shots)
    verified = bool(
        outcome
        and outcome.error is None
        and shots
        and not omitted
        and not spill_unknown
        and not cropped
        and (not calc_required or calc == "done")
    )
    calculated = calc == "done"
    ok = committed and verified and (not calc_required or calculated)
    if error:
        ok = False
        verified = False
    ranges = {shot.sheet: shot.range for shot in shots}
    suggestions = []
    if committed and not verified:
        suggestions.append({"tool": tool, "ops": [{"action": "verify", "ranges": ranges}]})
    body = {
        "ok": ok,
        "committed": committed,
        "saved": committed,
        "verified": verified,
        "calculated": calculated,
        "calc": calc,
        "engine": engine,
        "path": path,
        "completedOps": completed,
        "failedAt": failed_at,
        "wrote": wrote,
        "screenshots": [
            {
                "sheet": shot.sheet,
                "range": shot.range,
                "cropped": shot.cropped,
                "spillUnknown": shot.spill_unknown or spill_unknown,
            }
            for shot in shots
        ],
        "omitted": omitted,
        "suggestedNextActions": suggestions,
    }
    if error:
        body["error"] = error
    if outcome and outcome.error and "error" not in body:
        body["error"] = {"code": "verify_failed", "message": outcome.error}
    return body, shots


def _track(ops: list[dict]) -> tuple[list[dict], list[dict], list[str], dict[str, str]]:
    wrote: list[dict] = []
    formulas: list[dict] = []
    touched: list[str] = []
    renamed: dict[str, str] = {}

    def remember(sheet: str) -> None:
        if sheet and sheet not in touched:
            touched.append(sheet)

    for op in ops:
        action = op["action"]
        sheet = str(op.get("sheet") or "")
        if action in {"insert_rows", "delete_rows", "insert_columns", "delete_columns"} and sheet:
            for item in wrote:
                if item["sheet"] == sheet:
                    item["range"] = shift_address(item["range"], op)
            for item in formulas:
                if item["sheet"] == sheet:
                    item["cell"] = shift_address(item["cell"], op).split(":")[0]
        if action == "rename_sheet":
            old = str(op.get("oldName") or sheet)
            new = str(op.get("newName"))
            renamed[new] = old
            for item in wrote:
                if item["sheet"] == old:
                    item["sheet"] = new
            for item in formulas:
                if item["sheet"] == old:
                    item["sheet"] = new
            if old in touched:
                touched[touched.index(old)] = new
            else:
                remember(new)
            continue
        if action == "set_values":
            address = _block_range(str(op.get("range")), op.get("values") or [])
            wrote.append({"sheet": sheet, "range": address})
            remember(sheet)
            for cell, formula in _formula_cells(sheet, str(op.get("range")), op.get("values") or []):
                formulas.append({"sheet": sheet, "cell": cell, "formula": formula})
        elif action == "set_formulas":
            address = _block_range(str(op.get("range")), op.get("formulas") or [])
            wrote.append({"sheet": sheet, "range": address})
            remember(sheet)
            for cell, formula in _formula_cells(sheet, str(op.get("range")), op.get("formulas") or []):
                formulas.append({"sheet": sheet, "cell": cell, "formula": formula})
        elif op.get("range") and sheet:
            wrote.append({"sheet": sheet, "range": str(op["range"])})
            remember(sheet)
        elif action == "create_sheet":
            remember(str(op.get("name") or sheet))
        elif action in {"copy_sheet", "delete_sheet"}:
            remember(str(op.get("newName") or op.get("source") or sheet))
        elif action in {"insert_rows", "delete_rows"}:
            row = int(op["row"])
            count = int(op.get("count") or 1)
            wrote.append({"sheet": sheet, "range": f"A{row}:A{row + count - 1}"})
            remember(sheet)
        elif action in {"insert_columns", "delete_columns"}:
            wrote.append({"sheet": sheet, "range": f"{op['column']}1"})
            remember(sheet)
    return wrote, formulas, touched, renamed


def _block_range(start: str, rows: list) -> str:
    from openpyxl.utils import get_column_letter, range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(start)
    height = max(len(rows), 1)
    width = max((len(row) for row in rows if isinstance(row, list)), default=1)
    end_row = max(max_row, min_row + height - 1)
    end_col = max(max_col, min_col + width - 1)
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(end_col)}{end_row}"


def _formula_cells(sheet: str, start: str, rows: list) -> list[tuple[str, str]]:
    from openpyxl.utils import get_column_letter, range_boundaries

    min_col, min_row, _, _ = range_boundaries(start)
    found = []
    for r_index, row in enumerate(rows):
        if not isinstance(row, list):
            continue
        for c_index, value in enumerate(row):
            if isinstance(value, str) and value.startswith("="):
                found.append((f"{get_column_letter(min_col + c_index)}{min_row + r_index}", value))
    del sheet
    return found


def _run_verify(
    session: ComSession,
    path: str,
    *,
    calc_required: bool,
    wrote: list[dict],
    formula_cells: list[dict],
    touched: list[str],
    renamed: dict[str, str],
    explicit_ranges: dict[str, str] | None = None,
) -> VerifyOutcome:
    outcome = VerifyOutcome(calc="skipped", calculated=False)
    try:
        if calc_required:
            calculate(session)
            outcome.calc = "done"
            outcome.calculated = True
        ranges, omitted, _ = plan_sheets(path, wrote, formula_cells, touched, renamed)
        outcome.omitted = omitted
        for sheet, pieces in _object_anchors(session).items():
            if touched and sheet not in touched and sheet not in ranges:
                continue
            merged = union_ranges([ranges.get(sheet, ""), *pieces])
            if merged:
                ranges[sheet] = merged
        if explicit_ranges:
            ranges.update(explicit_ranges)
        spill_unknown = False
        for item in formula_cells:
            supported, address = spill_address(session, item["sheet"], item["cell"])
            if not supported:
                spill_unknown = True
                continue
            if address:
                ranges[item["sheet"]] = union_ranges([ranges.get(item["sheet"], address), address]) or address
        outcome.spill_unknown = spill_unknown
        if not ranges and touched:
            for sheet in touched:
                shot = capture_range(session, sheet, "A1:L40")
                shot.cropped = True
                shot.spill_unknown = spill_unknown
                outcome.shots.append(shot)
        for sheet, address in ranges.items():
            if not address:
                continue
            shot = capture_range(session, sheet, address)
            shot.spill_unknown = spill_unknown
            outcome.shots.append(shot)
        if spill_unknown or omitted or any(shot.cropped for shot in outcome.shots):
            outcome.error = outcome.error
        reasons = []
        if spill_unknown:
            reasons.append("spill range unavailable")
        if omitted:
            reasons.append("dependent ranges omitted")
        if any(shot.cropped for shot in outcome.shots):
            reasons.append("screenshot cropped")
        if not outcome.shots:
            reasons.append("no screenshot range")
        if reasons:
            outcome.error = "; ".join(reasons)
    except Exception as exc:
        if calc_required and outcome.calc != "done":
            outcome.calc = "failed"
            outcome.calculated = False
            outcome.error = str(exc)
        else:
            outcome.error = str(exc)
    return outcome


def apply_workbook(path: str, ops: list[dict], tool: str = "workbook_apply") -> tuple[dict, list]:
    try:
        ops = expand_layouts(normalize_ops(ops))
    except ValueError as exc:
        return _error("invalid_ops", str(exc), path=path), []
    unknown = validate_known(ops)
    if unknown:
        return _error("unsupported", f"unsupported action: {unknown}", path=path), []
    if any(op["action"] == "verify" for op in ops) and len(ops) != 1:
        return _error("verify_must_be_alone", "verify cannot be combined with other ops", path=path), []
    if count_cells(ops) > MAX_WRITE_CELLS:
        return _error("too_large", f"a single call can write at most {MAX_WRITE_CELLS} cells", path=path), []
    path = canonical_path(path)
    bad = _extension(path)
    if bad:
        return _error("unsupported_format", f"only .xlsx and .xlsm are accepted, got {bad}", path=path), []

    creating = not os.path.exists(path)
    if creating:
        if not ops or ops[0]["action"] != "create_workbook":
            return _error("not_found", "workbook does not exist", path=path), []
        create_empty(path)
        ops = [op for op in ops if op["action"] != "create_workbook"]
        if not ops:
            ops = [{"action": "verify", "ranges": {"Sheet1": "A1"}}]
    elif is_encrypted_or_irm(path):
        return _error("encrypted_or_irm", "encrypted or IRM workbooks are rejected before modification", path=path), []

    try:
        lock = FileLock(path)
        lock.acquire()
    except LockTimeout:
        return _error("lock_timeout", "timed out waiting for the workbook lock", path=path), []

    try:
        if file_is_locked(path):
            return _error("file_locked", "workbook is open in Excel or another process", path=path), []
        info = inspect_workbook(path)
        if ops and ops[0]["action"] == "verify" and len(ops) == 1:
            decision = RouteDecision("com", True, False, "verify")
        else:
            decision = decide(ops, info)
        if decision.engine == "openpyxl":
            body, shots = _apply_openpyxl(path, ops, tool, decision, lock)
        else:
            body, shots = _apply_com(path, ops, tool, decision, lock)
        return body, shots
    finally:
        lock.release()


def _apply_openpyxl(path: str, ops: list[dict], tool: str, decision: RouteDecision, lock: FileLock):
    workbook = load_for_edit(path)
    applied = apply_ops(workbook, ops)
    if applied.error:
        workbook.close()
        return _response(
            path=path,
            engine="openpyxl",
            tool=tool,
            committed=False,
            calc="skipped",
            calc_required=decision.calc_required,
            wrote=applied.wrote,
            outcome=None,
            completed=applied.completed,
            failed_at=applied.failed_at,
            error={"code": "op_failed", "message": applied.error},
        )
    temp_path = None
    try:
        temp_path = save_atomic(workbook, path)
    except Exception as exc:
        return _response(
            path=path,
            engine="openpyxl",
            tool=tool,
            committed=False,
            calc="failed" if decision.calc_required else "skipped",
            calc_required=decision.calc_required,
            wrote=applied.wrote,
            outcome=None,
            completed=applied.completed,
            failed_at=applied.failed_at,
            error={"code": "save_failed", "message": str(exc)},
        )
    finally:
        workbook.close()

    outcome = VerifyOutcome(calc="skipped", calculated=False)
    committed = False
    try:
        if decision.calc_required:
            with COM_QUEUE:
                session = ComSession(temp_path, allow_macros=False)
                try:
                    session.open()
                    outcome = _run_verify(
                        session,
                        temp_path,
                        calc_required=True,
                        wrote=applied.wrote,
                        formula_cells=applied.formula_cells,
                        touched=applied.touched_sheets,
                        renamed={},
                    )
                    if outcome.calc != "done":
                        freed = session.abort()
                        if not freed:
                            lock.retain_until_file_free()
                        discard(temp_path)
                        return _response(
                            path=path,
                            engine="openpyxl",
                            tool=tool,
                            committed=False,
                            calc="failed",
                            calc_required=True,
                            wrote=applied.wrote,
                            outcome=outcome,
                            completed=applied.completed,
                            failed_at=applied.failed_at,
                            error={"code": "calc_failed", "message": outcome.error or "calculation failed"},
                        )
                    session.close(save=True)
                except Exception as exc:
                    freed = session.abort()
                    if not freed:
                        lock.retain_until_file_free()
                    discard(temp_path)
                    return _response(
                        path=path,
                        engine="openpyxl",
                        tool=tool,
                        committed=False,
                        calc="failed",
                        calc_required=True,
                        wrote=applied.wrote,
                        outcome=None,
                        completed=applied.completed,
                        failed_at=applied.failed_at,
                        error={"code": "excel_busy" if not session.pid_known else "calc_failed", "message": str(exc)},
                    )
            replace_target(temp_path, path)
            committed = True
            temp_path = None
        else:
            replace_target(temp_path, path)
            committed = True
            temp_path = None
            with COM_QUEUE:
                session = ComSession(path, allow_macros=False)
                try:
                    session.open()
                    outcome = _run_verify(
                        session,
                        path,
                        calc_required=False,
                        wrote=applied.wrote,
                        formula_cells=[],
                        touched=applied.touched_sheets,
                        renamed={},
                    )
                    session.close(save=False)
                except Exception as exc:
                    freed = session.abort()
                    if not freed:
                        lock.retain_until_file_free()
                    outcome = VerifyOutcome(calc="skipped", calculated=False, error=str(exc))
    except Exception as exc:
        discard(temp_path)
        return _response(
            path=path,
            engine="openpyxl",
            tool=tool,
            committed=committed,
            calc=outcome.calc,
            calc_required=decision.calc_required,
            wrote=applied.wrote,
            outcome=outcome,
            completed=applied.completed,
            failed_at=applied.failed_at,
            error={"code": "save_failed", "message": str(exc), "tempPath": temp_path},
        )
    error = None
    if outcome.error:
        error = {"code": "verify_failed", "message": outcome.error}
    return _response(
        path=path,
        engine="openpyxl",
        tool=tool,
        committed=committed,
        calc=outcome.calc,
        calc_required=decision.calc_required,
        wrote=applied.wrote,
        outcome=outcome,
        completed=applied.completed,
        failed_at=None,
        error=error,
    )


def _apply_com(path: str, ops: list[dict], tool: str, decision: RouteDecision, lock: FileLock):
    if ops and ops[0]["action"] == "verify" and len(ops) == 1:
        return _verify_only(path, ops[0], tool, lock)
    wrote, formulas, touched, renamed = _track(ops)
    commands = []
    for op in ops:
        if op["action"] == "create_workbook":
            continue
        command = COM_COMMANDS.get(op["action"])
        if not command:
            return _response(
                path=path,
                engine="com",
                tool=tool,
                committed=False,
                calc="skipped",
                calc_required=decision.calc_required,
                wrote=[],
                outcome=None,
                completed=0,
                failed_at=0,
                error={"code": "unsupported", "message": op["action"]},
            ), []
        commands.append({"command": command, "args": com_args(op)})
    slow = {"table_add_to_data_model", "table_create_from_dax"}
    timeout = 600 if any(op["action"].startswith(("powerquery", "datamodel")) or op["action"] in slow for op in ops) else None
    completed = 0
    with COM_QUEUE:
        session = ComSession(path, allow_macros=decision.allow_macros, timeout=timeout)
        try:
            session.open()
        except ExcelCliError as exc:
            freed = session.abort()
            if not freed:
                lock.retain_until_file_free()
            code = "excel_busy" if not exc.pid_known else "op_failed"
            return _response(
                path=path,
                engine="com",
                tool=tool,
                committed=False,
                calc="skipped",
                calc_required=decision.calc_required,
                wrote=[],
                outcome=None,
                completed=0,
                failed_at=0,
                error={"code": code, "message": str(exc)},
            )
        for index, command in enumerate(commands):
            try:
                session.call([command])
                completed = index + 1
            except ExcelCliError as exc:
                freed = session.abort()
                if not freed:
                    lock.retain_until_file_free()
                return _response(
                    path=path,
                    engine="com",
                    tool=tool,
                    committed=False,
                    calc="skipped",
                    calc_required=decision.calc_required,
                    wrote=wrote,
                    outcome=None,
                    completed=completed,
                    failed_at=index,
                    error={"code": "op_failed", "message": str(exc)},
                )
        outcome = _run_verify(
            session,
            path,
            calc_required=decision.calc_required,
            wrote=wrote,
            formula_cells=formulas,
            touched=touched,
            renamed=renamed,
        )
        if decision.calc_required and outcome.calc != "done":
            freed = session.abort()
            if not freed:
                lock.retain_until_file_free()
            return _response(
                path=path,
                engine="com",
                tool=tool,
                committed=False,
                calc="failed",
                calc_required=True,
                wrote=wrote,
                outcome=outcome,
                completed=completed,
                failed_at=None,
                error={"code": "calc_failed", "message": outcome.error or "calculation failed"},
            )
        try:
            session.close(save=True)
        except ExcelCliError as exc:
            freed = session.abort()
            if not freed:
                lock.retain_until_file_free()
            return _response(
                path=path,
                engine="com",
                tool=tool,
                committed=False,
                calc=outcome.calc,
                calc_required=decision.calc_required,
                wrote=wrote,
                outcome=outcome,
                completed=completed,
                failed_at=None,
                error={"code": "save_failed" if session.pid_known else "excel_busy", "message": str(exc)},
            )
    error = {"code": "verify_failed", "message": outcome.error} if outcome.error else None
    if decision.calc_required and outcome.calc != "done":
        error = error or {"code": "calc_failed", "message": outcome.error or "calculation failed"}
    return _response(
        path=path,
        engine="com",
        tool=tool,
        committed=True,
        calc=outcome.calc,
        calc_required=decision.calc_required,
        wrote=wrote,
        outcome=outcome,
        completed=completed,
        failed_at=None,
        error=error,
    )


def _verify_only(path: str, op: dict, tool: str, lock: FileLock):
    explicit = op.get("ranges") or {}
    sheets = list(op.get("sheets") or explicit.keys())
    wrote = [{"sheet": sheet, "range": explicit.get(sheet, "A1:L40")} for sheet in sheets]
    with COM_QUEUE:
        session = ComSession(path, allow_macros=False)
        try:
            session.open()
            outcome = _run_verify(
                session,
                path,
                calc_required=True,
                wrote=wrote,
                formula_cells=[],
                touched=sheets,
                renamed={},
                explicit_ranges=explicit or None,
            )
            if outcome.calc != "done":
                freed = session.abort()
                if not freed:
                    lock.retain_until_file_free()
                return _response(
                    path=path,
                    engine="com",
                    tool=tool,
                    committed=False,
                    calc="failed",
                    calc_required=True,
                    wrote=wrote,
                    outcome=outcome,
                    completed=0,
                    failed_at=0,
                    error={"code": "calc_failed", "message": outcome.error or "calculation failed"},
                )
            session.close(save=True)
        except ExcelCliError as exc:
            freed = session.abort()
            if not freed:
                lock.retain_until_file_free()
            return _response(
                path=path,
                engine="com",
                tool=tool,
                committed=False,
                calc="failed",
                calc_required=True,
                wrote=wrote,
                outcome=None,
                completed=0,
                failed_at=0,
                error={"code": "verify_failed", "message": str(exc)},
            )
    error = {"code": "verify_failed", "message": outcome.error} if outcome.error else None
    return _response(
        path=path,
        engine="com",
        tool=tool,
        committed=True,
        calc=outcome.calc,
        calc_required=True,
        wrote=wrote,
        outcome=outcome,
        completed=1,
        failed_at=None,
        error=error,
    )


def _with_images(body: dict, shots: list) -> list:
    content: list[Any] = [body]
    for shot in shots:
        if not getattr(shot, "data", None):
            continue
        fmt = "png" if "png" in shot.mime else "jpeg"
        content.append(Image(data=shot.data, format=fmt))
    return content


@mcp.tool(structured_output=False)
def workbook_read(
    path: str,
    sheet: str | None = None,
    range: str | None = None,
    includeStyles: bool = False,
    limit: int = READ_LIMIT,
) -> list:
    """Read values, formulas, and cached values. Takes the workbook lock. Does not screenshot.

    Pass the same ops shape you would write back when includeStyles is true.
    Results are paged. Follow nextRange until it is null. Default page is 4000 cells.
    Do not read a file that is open in Excel; close it first.
    """
    full = canonical_path(path)
    bad = _extension(full)
    if bad:
        return [_error("unsupported_format", f"only .xlsx and .xlsm are accepted, got {bad}", path=full)]
    if not os.path.exists(full):
        return [_error("not_found", "workbook does not exist", path=full)]
    if is_encrypted_or_irm(full):
        return [_error("encrypted_or_irm", "encrypted or IRM workbooks are rejected", path=full)]
    try:
        with FileLock(full):
            if file_is_locked(full):
                return [_error("file_locked", "workbook is open in Excel or another process", path=full)]
            page = read_page(full, sheet, range, includeStyles, min(limit, READ_LIMIT))
    except LockTimeout:
        return [_error("lock_timeout", "timed out waiting for the workbook lock", path=full)]
    except Exception as exc:
        return [_error("op_failed", str(exc), path=full)]
    page["ok"] = True
    page["path"] = full
    return [page]


@mcp.tool(structured_output=False)
def workbook_apply(path: str, ops: list[dict[str, Any]]) -> list:
    """Apply a batch of ops, save once, recalculate when required, then screenshot.

    Do not pass engine or session_id. Structural ops, VBA, tables, pivots, charts,
    Power Query, and data-model ops are routed to Excel automatically.
    Put every change to this file in one ops list. After committed is true, do not
    send the same structural ops again. If verified is false, call verify only.
    Formulas use English function names. Number formats use Excel format codes.
    Close the file in Excel before calling. Screenshot is included; do not ask for
    another image. Do not paint Excel table bodies with format.
    For a new sheet, send action layout. Profiles share one palette; blocks differ by domain.
    """
    body, shots = apply_workbook(path, ops, "workbook_apply")
    return _with_images(body, shots)


@mcp.tool(structured_output=False)
def excel_exec(path: str, ops: list[dict[str, Any]]) -> list:
    """Same ops as workbook_apply, including VBA, tables, pivots, charts, Power Query, and DAX.

    The server still chooses the engine. VBA runs only when an op in this call
    imports, updates, or runs VBA, and Excel must trust access to the VBA project.
    Screenshot is part of this call. A failed verify is retried with action verify,
    never by repeating an insert or delete.
    """
    body, shots = apply_workbook(path, ops, "excel_exec")
    return _with_images(body, shots)


def main() -> None:
    mcp.run()


if __name__ == "__main__":
    main()
