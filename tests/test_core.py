"""Routing, lock, package scan, and commit-result tests that do not need Excel."""

from __future__ import annotations

import os
import threading
import zipfile
from pathlib import Path

from openpyxl import Workbook

from excel_sovereign.book import apply_ops, file_hash, load_for_edit, save_atomic
from excel_sovereign.inspect import inspect_workbook
from excel_sovereign.lock import FileLock, canonical_path
from excel_sovereign.route import KNOWN_ACTIONS, TOOL_ACTIONS, action_outside, com_args, count_cells, decide, normalize_ops
from excel_sovereign.server import _response, apply_workbook
from excel_sovereign.verify import VerifyOutcome


def _xlsx(path: Path, value=1) -> None:
    workbook = Workbook()
    workbook.active.title = "参数"
    workbook.active["A1"] = value
    workbook.save(path)
    workbook.close()


def test_plain_style_stays_on_openpyxl(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _xlsx(path)
    info = inspect_workbook(str(path))
    ops = normalize_ops([{"action": "format", "sheet": "参数", "range": "A1", "fontColor": "#FF0000"}])
    decision = decide(ops, info)
    assert decision.engine == "openpyxl"
    assert decision.calc_required is False


def test_com_part_promotes_font_change(tmp_path: Path):
    path = tmp_path / "pivot.xlsx"
    _xlsx(path)
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("xl/pivotCaches/pivotCacheDefinition1.xml", "<pivotCacheDefinition/>")
    info = inspect_workbook(str(path))
    ops = normalize_ops([{"action": "format", "sheet": "参数", "range": "A1", "fontColor": "#FF0000"}])
    decision = decide(ops, info)
    assert decision.engine == "com"
    assert "xl/pivotCaches/pivotCacheDefinition1.xml" in info.com_parts


def test_rename_and_insert_require_com(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _xlsx(path)
    info = inspect_workbook(str(path))
    rename = normalize_ops([{"action": "rename_sheet", "oldName": "参数", "newName": "数据"}])
    insert = normalize_ops([{"action": "insert_rows", "sheet": "参数", "row": 2, "count": 1}])
    assert decide(rename, info).engine == "com"
    assert decide(insert, info).engine == "com"


def test_value_write_recalculates_when_formula_exists(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = 1
    workbook.active["B1"] = "=A1+1"
    workbook.save(path)
    workbook.close()
    info = inspect_workbook(str(path))
    ops = normalize_ops([{"action": "set_values", "sheet": "Sheet", "range": "A1", "values": [[5]]}])
    decision = decide(ops, info)
    assert decision.calc_required is True


def test_table_intersection_promotes(tmp_path: Path):
    path = tmp_path / "table.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "n"
    sheet["A2"] = 1
    from openpyxl.worksheet.table import Table, TableStyleInfo

    table = Table(displayName="Data", ref="A1:A2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    workbook.save(path)
    workbook.close()
    info = inspect_workbook(str(path))
    ops = normalize_ops([{"action": "set_values", "sheet": sheet.title, "range": "A2", "values": [[3]]}])
    assert decide(ops, info).engine == "com"


def test_too_large_and_unknown_action(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _xlsx(path)
    huge = [{"action": "set_values", "sheet": "参数", "range": "A1", "values": [[1] for _ in range(100_001)]}]
    body, _ = apply_workbook(str(path), huge)
    assert body["error"]["code"] == "too_large"
    assert body["committed"] is False
    body, _ = apply_workbook(str(path), [{"action": "not_a_real_action"}])
    assert body["error"]["code"] == "unsupported"


def test_encrypted_ole_rejected(tmp_path: Path):
    path = tmp_path / "secret.xlsx"
    path.write_bytes(bytes([0xD0, 0xCF, 0x11, 0xE0, 0xA1, 0xB1, 0x1A, 0xE1]) + b"\x00" * 16)
    body, _ = apply_workbook(str(path), [{"action": "set_values", "sheet": "S", "range": "A1", "values": [[1]]}])
    assert body["error"]["code"] == "encrypted_or_irm"
    assert body["committed"] is False


def test_verify_suggestion_does_not_repeat_insert():
    outcome = VerifyOutcome(calc="done", calculated=True, error="screenshot timed out")
    body, _ = _response(
        path="D:/books/a.xlsx",
        engine="com",
        tool="workbook_apply",
        committed=True,
        calc="done",
        calc_required=True,
        wrote=[{"sheet": "参数", "range": "A2:A2"}],
        outcome=outcome,
        completed=1,
        failed_at=None,
        error={"code": "verify_failed", "message": "screenshot timed out"},
    )
    assert body["committed"] is True
    assert body["verified"] is False
    assert body["ok"] is False
    actions = body["suggestedNextActions"][0]["ops"]
    assert actions == [{"action": "verify", "ranges": {}}]
    assert all(op["action"] != "insert_rows" for op in actions)


def test_lock_blocks_reader_until_release(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _xlsx(path)
    started = threading.Event()
    release = threading.Event()
    order = []

    def holder():
        with FileLock(str(path)):
            started.set()
            release.wait(5)
            order.append("released")

    thread = threading.Thread(target=holder)
    thread.start()
    assert started.wait(5)
    second = FileLock(str(path), timeout=5)
    acquired = threading.Event()

    def waiter():
        second.acquire()
        order.append("acquired")
        second.release()
        acquired.set()

    waiter_thread = threading.Thread(target=waiter)
    waiter_thread.start()
    assert not acquired.wait(0.4)
    release.set()
    assert acquired.wait(5)
    thread.join(5)
    waiter_thread.join(5)
    assert order == ["released", "acquired"]


def test_xlsm_keeps_vba_project(tmp_path: Path):
    path = tmp_path / "macro.xlsm"
    workbook = Workbook()
    workbook.active["A1"] = "x"
    workbook.save(path)
    workbook.close()
    marker = b"vba-marker-not-a-real-project"
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("xl/vbaProject.bin", marker)
    loaded = load_for_edit(str(path))
    apply_ops(loaded, normalize_ops([{"action": "format", "sheet": "Sheet", "range": "A1", "fontColor": "#00FF00"}]))
    temp = save_atomic(loaded, str(path))
    loaded.close()
    os.replace(temp, path)
    with zipfile.ZipFile(path) as archive:
        assert archive.read("xl/vbaProject.bin") == marker
    info = inspect_workbook(str(path))
    assert info.has_vba is True
    assert info.has_com_parts is False


def test_failed_openpyxl_op_does_not_change_bytes(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _xlsx(path, value=7)
    before = file_hash(str(path))
    body, _ = apply_workbook(
        str(path),
        [{"action": "set_values", "sheet": "missing", "range": "A1", "values": [[1]]}],
    )
    # This reaches Excel for screenshot only if the sheet exists. A missing sheet fails first.
    assert body["committed"] is False
    assert file_hash(str(path)) == before


def test_cell_counter_matches_values():
    ops = normalize_ops([{"action": "set_values", "sheet": "S", "range": "A1", "values": [[1, 2], [3, 4]]}])
    assert count_cells(ops) == 4


def test_canonical_path_folds_case():
    assert canonical_path("C:/Temp/A.xlsx") == canonical_path("c:/temp/a.xlsx")


def test_actions_belong_to_one_tool():
    owners = {}
    for tool, actions in TOOL_ACTIONS.items():
        assert "verify" in actions
        for action in actions:
            if action == "verify":
                continue
            assert action not in owners, action
            owners[action] = tool
    assert KNOWN_ACTIONS | {"layout"} <= set(owners) | {"verify"}


def test_other_family_is_rejected_before_write(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _xlsx(path)
    before = file_hash(str(path))
    body, _ = apply_workbook(
        str(path),
        [{"action": "pivot_create_from_range", "sheet": "参数", "range": "A1:B2"}],
        "workbook_apply",
    )
    assert body["error"]["code"] == "wrong_tool"
    assert body["committed"] is False
    assert body["suggestedNextActions"] == [
        {"tool": "excel_table", "ops": [{"action": "pivot_create_from_range"}]}
    ]
    assert file_hash(str(path)) == before
    assert action_outside("excel_vba", [{"action": "vba_run", "macro": "Module1.Main"}]) is None


def test_com_write_range_matches_matrix():
    args = com_args({"action": "set_values", "sheet": "Sheet1", "range": "A1", "values": [["a", "b"], [1, 2]]})
    assert args["rangeAddress"] == "A1:B2"
