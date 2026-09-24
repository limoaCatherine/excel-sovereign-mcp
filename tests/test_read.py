"""Read modes, paging, caching, and compact output. openpyxl only."""

from __future__ import annotations

import json
import os
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from excel_sovereign.server import read_workbook, workbook_read


def _book(path: Path) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "技能"
    ws["A1"] = "技能总表"
    ws.merge_cells("A1:D1")
    ws["A5"] = "技能ID"
    ws["B5"] = "名称"
    ws["D5"] = "系数"
    ws["A9"] = 1001
    ws["B9"] = "火球术" * 30
    ws["D9"] = 1.5
    ws["E9"] = "=D9*2"
    ws["Z400"].font = Font(bold=True)
    ws.freeze_panes = "B6"
    other = workbook.create_sheet("空表")
    other.sheet_state = "hidden"
    workbook.save(path)
    workbook.close()


def test_sparse_returns_only_filled_cells(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    page = read_workbook(str(path), sheet="技能")
    assert page["ok"] is True
    assert page["range"] == "A1:E9"
    assert list(page["rows"]) == ["1", "5", "9"]
    assert page["rows"]["5"] == {"A": "技能ID", "B": "名称", "D": "系数"}
    assert page["rows"]["9"]["E"] == {"f": "=D9*2", "v": None}
    assert page["merged"] == ["A1:D1"]
    assert page["nextRange"] is None


def test_sparse_pages_on_row_boundary(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    first = read_workbook(str(path), sheet="技能", limit=2)
    assert list(first["rows"]) == ["1", "5"]
    assert first["nextRange"] == "A9:E9"
    second = read_workbook(str(path), sheet="技能", range=first["nextRange"], limit=2)
    assert list(second["rows"]) == ["9"]
    assert second["nextRange"] is None


def test_dense_keeps_write_shape_and_drops_empty_formulas(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    plain = read_workbook(str(path), sheet="技能", range="A5:D5", mode="dense")
    assert plain["values"] == [["技能ID", "名称", None, "系数"]]
    assert "formulas" not in plain
    mixed = read_workbook(str(path), sheet="技能", range="D9:E9", mode="dense")
    assert mixed["formulas"] == [[None, "=D9*2"]]


def test_dense_pages_mid_row(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    page = read_workbook(str(path), sheet="技能", range="A5:D6", mode="dense", limit=6)
    assert page["range"] == "A5:B6"
    assert page["nextRange"] == "C6:D6"


def test_overview_lists_every_sheet(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    page = read_workbook(str(path), mode="overview", preview=2)
    skill, empty = page["sheets"]
    assert skill["name"] == "技能"
    assert skill["usedRange"] == "A1:E9"
    assert skill["cells"] == 8
    assert skill["formulas"] == 1
    assert skill["merged"] == 1
    assert skill["freeze"] == "B6"
    assert list(skill["preview"]) == ["1", "5"]
    assert empty == {"name": "空表", "usedRange": None, "cells": 0, "formulas": 0, "merged": 0, "state": "hidden"}


def test_max_text_clips_strings(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    page = read_workbook(str(path), sheet="技能", range="B9", maxText=6)
    assert page["rows"]["9"]["B"] == "火球术火球术…"


def test_styles_ride_on_sparse_cells(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    page = read_workbook(str(path), sheet="技能", range="A5", includeStyles=True)
    cell = page["rows"]["5"]["A"]
    assert cell["v"] == "技能ID"
    assert "fontName" in cell["s"]


def test_read_does_not_grow_sheet_and_cache_sees_writes(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    read_workbook(str(path), sheet="技能", range="A1:H20", mode="dense")
    again = read_workbook(str(path), mode="overview")
    assert again["sheets"][0]["usedRange"] == "A1:E9"
    workbook = load_workbook(path)
    workbook["技能"]["F9"] = "新列"
    workbook.save(path)
    workbook.close()
    stat = os.stat(path)
    os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000))
    fresh = read_workbook(str(path), sheet="技能", range="F9")
    assert fresh["rows"]["9"]["F"] == "新列"


def test_bad_mode_and_missing_sheet_are_errors(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    assert read_workbook(str(path), mode="grid")["error"]["code"] == "op_failed"
    assert read_workbook(str(path), sheet="不存在")["ok"] is False


def test_empty_sheet_page(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    page = read_workbook(str(path), sheet="空表")
    assert page["empty"] is True
    assert page["nextRange"] is None


def test_find_searches_values_and_formulas(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    page = read_workbook(str(path), mode="find", find="d9")
    assert page["matches"] == [{"sheet": "技能", "cell": "E9", "f": "=D9*2"}]
    names = read_workbook(str(path), mode="find", find="技能")
    assert [hit["cell"] for hit in names["matches"]] == ["A1", "A5"]
    capped = read_workbook(str(path), mode="find", find="技能", limit=1)
    assert capped["truncated"] is True
    scoped = read_workbook(str(path), mode="find", find="技能", range="A5:B9")
    assert [hit["cell"] for hit in scoped["matches"]] == ["A5"]
    assert read_workbook(str(path), mode="find")["ok"] is False


def test_overview_flags_formatting_past_data_and_trim_plans_it(tmp_path: Path):
    from excel_sovereign.book import trim_ops

    path = tmp_path / "a.xlsx"
    _book(path)
    workbook = load_workbook(path)
    workbook["技能"].column_dimensions["AB"].width = 30
    workbook.save(path)
    workbook.close()
    skill = read_workbook(str(path), mode="overview")["sheets"][0]
    assert skill["extent"] == "A1:AB400"
    assert trim_ops(str(path), "技能") == [
        {"action": "delete_columns", "sheet": "技能", "column": "F", "count": 23},
        {"action": "delete_rows", "sheet": "技能", "row": 10, "count": 391},
    ]
    clean = tmp_path / "clean.xlsx"
    workbook = Workbook()
    workbook.active["B2"] = 1
    workbook.save(clean)
    workbook.close()
    assert trim_ops(str(clean), "Sheet") == []
    assert "extent" not in read_workbook(str(clean), mode="overview")["sheets"][0]
    workbook = load_workbook(clean)
    tail = workbook.active.column_dimensions["C"]
    tail.min, tail.max, tail.width = 3, 16384, 9
    workbook.save(clean)
    workbook.close()
    assert trim_ops(str(clean), "Sheet") == []


def test_trim_keeps_what_formulas_names_and_validations_point_at(tmp_path: Path):
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    from excel_sovereign.book import trim_ops

    path = tmp_path / "refs.xlsx"
    workbook = Workbook()
    data = workbook.active
    data.title = "数据"
    data["A1"] = 1
    data["Z300"].font = Font(bold=True)
    other = workbook.create_sheet("汇总")
    other["A1"] = "=SUM(数据!$A$1:$C$40)"
    other["A2"] = '="数据!Z999"'
    other["A3"] = "=LOG10(4)"
    workbook.defined_names["池"] = DefinedName("池", attr_text="数据!$A$1:$A$60")
    rule = DataValidation(type="whole")
    rule.add("E1:E20")
    data.add_data_validation(rule)
    workbook.save(path)
    workbook.close()
    assert trim_ops(str(path), "数据") == [
        {"action": "delete_columns", "sheet": "数据", "column": "F", "count": 21},
        {"action": "delete_rows", "sheet": "数据", "row": 61, "count": 240},
    ]


def test_trim_sheet_with_nothing_to_trim_is_a_no_op(tmp_path: Path):
    from excel_sovereign.book import file_hash
    from excel_sovereign.server import apply_workbook

    path = tmp_path / "clean.xlsx"
    workbook = Workbook()
    workbook.active["B2"] = 1
    workbook.save(path)
    workbook.close()
    before = file_hash(str(path))
    body, shots = apply_workbook(str(path), [{"action": "trim_sheet", "sheet": "Sheet"}])
    assert body["ok"] is True and body["committed"] is False and shots == []
    assert file_hash(str(path)) == before
    missing, _ = apply_workbook(str(path), [{"action": "trim_sheet", "sheet": "无"}])
    assert missing["error"]["code"] == "invalid_ops"


def test_read_works_while_another_process_holds_the_file(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    with open(path, "rb"):
        page = read_workbook(str(path), sheet="技能", range="A5")
    assert page["ok"] is True
    assert page["openInExcel"] is True
    assert page["rows"]["5"]["A"] == "技能ID"
    assert "openInExcel" not in read_workbook(str(path), sheet="技能", range="A5")


def test_query_results_map_back_to_caller_ops():
    from excel_sovereign.server import _results_of

    ops = [{"action": "create_workbook"}, {"action": "set_values"}, {"action": "table_list"}]
    results = [
        {"index": 0, "command": "range.set-values", "success": True, "result": {"success": True}},
        {"index": 1, "command": "table.list", "success": True, "result": {"tables": ["T"], "success": True, "filePath": "x"}},
    ]
    assert _results_of(ops, [1, 2], results) == [{"op": 2, "action": "table_list", "result": {"tables": ["T"]}}]


def test_track_reads_targets_inside_args_and_skips_gone_sheets():
    from excel_sovereign.server import _track

    wrote, _, touched, _ = _track(
        [
            {"action": "hyperlink_add", "args": {"sheetName": "S", "cellAddress": "A1"}},
            {"action": "create_sheet", "name": "Side"},
            {"action": "sheet_hide", "sheet": "Side"},
            {"action": "create_sheet", "name": "Gone"},
            {"action": "delete_sheet", "sheet": "Gone"},
            {"action": "sheet_show", "sheet": "Back"},
        ]
    )
    assert wrote == [{"sheet": "S", "range": "A1"}, {"sheet": "Back", "range": "A1:L40"}]
    assert touched == ["S", "Back"]


def test_excel_pids_is_a_set():
    from excel_sovereign.excel_cli import excel_pids

    assert isinstance(excel_pids(), set)


def test_tool_output_is_compact_json(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    _book(path)
    [text] = workbook_read(str(path), sheet="技能")
    assert isinstance(text, str)
    assert "\n" not in text
    assert "技能ID" in text
    assert json.loads(text)["rows"]["5"]["A"] == "技能ID"
