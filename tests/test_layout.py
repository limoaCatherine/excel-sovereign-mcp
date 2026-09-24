"""Layout tokens stay shared while finance, analytics, and other domains stack different blocks."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.utils import get_column_letter

from excel_sovereign.book import apply_ops, create_empty, load_for_edit, save_atomic
from excel_sovereign.layout import PROFILES, compile_layout
from excel_sovereign.server import apply_workbook


def _finance():
    return {
        "action": "layout",
        "sheet": "Sheet1",
        "profile": "finance",
        "blocks": [
            {"kind": "title", "text": "收入预测"},
            {
                "kind": "parameters",
                "columns": [
                    {"header": "项目", "role": "label"},
                    {"header": "值", "role": "percent", "input": True},
                    {"header": "单位", "role": "unit"},
                    {"header": "说明", "role": "text"},
                ],
                "rows": [
                    ["收入增速", 0.08, "%", "可调"],
                    ["税率", "=利润!C3", "%", "跨表"],
                    ["调整", "=SUM(科目[金额])", "元", "结构化引用"],
                    ["备注", None, "", "—"],
                ],
            },
            {
                "kind": "formula",
                "bands": [
                    {
                        "title": "利润",
                        "columns": [
                            {"header": "项目", "role": "label"},
                            {"header": "金额", "role": "number"},
                        ],
                        "rows": [["收入", 100]],
                    },
                    {
                        "title": "检查",
                        "columns": [{"header": "差额", "role": "formula"}],
                        "rows": [["=B12-C12"]],
                    },
                ],
            },
        ],
    }


def _catalog():
    return {
        "action": "layout",
        "sheet": "科目",
        "profile": "finance",
        "create": True,
        "blocks": [
            {"kind": "title", "text": "科目"},
            {
                "kind": "catalog",
                "group": "类别",
                "columns": [
                    {"header": "类别", "role": "group"},
                    {"header": "科目", "role": "label"},
                    {"header": "编码", "role": "id"},
                ],
                "rows": [["资产", "现金", 1001], ["资产", "应收", 1002], ["负债", "应付", 2001]],
            },
        ],
    }


def _analytics():
    return {
        "action": "layout",
        "sheet": "明细",
        "profile": "analytics",
        "create": True,
        "blocks": [
            {"kind": "title", "text": "订单明细"},
            {
                "kind": "fact",
                "name": "Orders",
                "columns": [
                    {"header": "区域", "role": "enum"},
                    {"header": "金额", "role": "number"},
                ],
                "rows": [["华东", 12], ["华北", 8]],
            },
        ],
    }


def _game():
    return {
        "action": "layout",
        "sheet": "经济",
        "profile": "general",
        "create": True,
        "blocks": [
            {"kind": "title", "text": "单局经济"},
            {
                "kind": "catalog",
                "group": "种类",
                "columns": [
                    {"header": "种类", "role": "group"},
                    {"header": "物品", "role": "label"},
                    {"header": "数量", "role": "number", "input": True},
                ],
                "rows": [["材料", "木材", 3], ["材料", "石料", 2], ["消耗", "口粮", 1]],
            },
        ],
    }


def _anchor(address: str):
    col, row, _, _ = range_boundaries(address)
    return col, row


def _cell_of(ops: list[dict], sheet: str, value):
    for op in ops:
        if op.get("action") != "set_values" or op.get("sheet") != sheet:
            continue
        col, row = _anchor(op["range"])
        for r_index, line in enumerate(op["values"]):
            for c_index, item in enumerate(line):
                if item == value:
                    return f"{get_column_letter(col + c_index)}{row + r_index}"
    raise AssertionError(value)


def _paint(ops: list[dict], sheet: str, address: str) -> dict:
    found = []
    for op in ops:
        if op.get("action") != "format" or op.get("sheet") != sheet:
            continue
        target = str(op.get("range") or "")
        if target == address or target.startswith(address + ":"):
            found.append(op)
    assert found, address
    return found[-1]


def test_profiles_do_not_invent_a_palette():
    finance = compile_layout(_finance())
    game = compile_layout(_game())
    assert _paint(finance, "Sheet1", "A1")["fontColor"] == _paint(game, "经济", "A1")["fontColor"]
    assert _paint(finance, "Sheet1", "A1")["fillColor"] == "#F2F2F2"
    assert "game" not in PROFILES
    assert "general" in PROFILES


def test_domains_stack_different_blocks():
    finance = compile_layout(_finance()) + compile_layout(_catalog())
    analytics = compile_layout(_analytics())
    assert not any(op["action"] == "table_create" for op in finance)
    table = next(op for op in analytics if op["action"] == "table_create")
    assert table["args"]["tableName"] == "Orders"
    fact_range = table["args"]["rangeAddress"]
    assert not any(op["action"] == "format" and op["range"] == fact_range for op in analytics)
    growth = _paint(finance, "Sheet1", _cell_of(finance, "Sheet1", 0.08))
    assert growth["fontColor"] == "#0000FF"
    assert growth["fillColor"] == "#F3F8FC"
    assert _paint(finance, "Sheet1", _cell_of(finance, "Sheet1", "=利润!C3"))["fontColor"] == "#008000"
    assert _paint(finance, "Sheet1", _cell_of(finance, "Sheet1", "=SUM(科目[金额])"))["fontColor"] == "#FF0000"
    payable = _cell_of(finance, "科目", "应付")
    row = "".join(ch for ch in payable if ch.isdigit())
    assert any(op.get("borderEdges") == "top" and op["range"].endswith(row) for op in finance if op["action"] == "format")
    section_fills = [op["range"] for op in finance if op.get("fillColor") == "#E7E6E6" and op["sheet"] == "Sheet1"]
    assert section_fills == ["A9:B9", "D9"]


def test_painted_sheet_uses_the_shared_tokens(tmp_path: Path):
    path = tmp_path / "finance.xlsx"
    create_empty(str(path))
    ops = compile_layout(_finance()) + compile_layout(_catalog())
    workbook = load_for_edit(str(path))
    applied = apply_ops(workbook, ops)
    assert applied.error is None
    temp = save_atomic(workbook, str(path))
    workbook.close()
    painted = load_workbook(temp)
    try:
        title = painted["Sheet1"]["A1"]
        assert title.value == "收入预测"
        assert title.font.color.rgb[-6:] == "1F4E79"
        assert title.font.bold is True
        assert title.fill.fgColor.rgb[-6:] == "F2F2F2"
        growth = painted["Sheet1"]["B4"]
        assert growth.value == 0.08
        assert growth.font.color.rgb[-6:] == "0000FF"
        assert growth.fill.fgColor.rgb[-6:] == "F3F8FC"
        assert painted["Sheet1"]["B5"].font.color.rgb[-6:] == "008000"
        assert painted["Sheet1"]["B6"].font.color.rgb[-6:] == "FF0000"
        assert painted["科目"]["A6"].border.top.style == "thin"
        assert painted["科目"]["A4"].border.top.style is None
    finally:
        painted.close()


def test_unknown_profile_is_rejected():
    body, _ = apply_workbook(
        r"C:\missing\none.xlsx",
        [{"action": "layout", "sheet": "S", "profile": "game", "blocks": [{"kind": "title", "text": "x"}]}],
    )
    assert body["error"]["code"] == "invalid_ops"
