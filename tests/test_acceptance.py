"""Acceptance checks that drive Excel. Skipped only when excelcli or Excel is missing."""

from __future__ import annotations

import json
import threading
import time
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Font

from excel_sovereign.book import file_hash
from excel_sovereign.excel_cli import ExcelCliError, excelcli_path, run_batch
from excel_sovereign.inspect import inspect_workbook
from excel_sovereign.lock import file_is_locked
from excel_sovereign.server import apply_workbook, read_workbook

pytestmark = pytest.mark.acceptance


def _require_excel() -> None:
    try:
        excelcli_path()
    except FileNotFoundError as exc:
        pytest.skip(str(exc))


def _apply(path: Path, ops: list[dict], tool: str = "workbook_apply"):
    body, shots = apply_workbook(str(path), ops, tool)
    return body, shots


def _create(path: Path, macro: bool = False) -> None:
    _require_excel()
    results = run_batch(
        [
            {
                "command": "session.create",
                "args": {
                    "filePath": str(path),
                    "macroEnabled": macro,
                    "timeoutSeconds": 120,
                },
            },
            {"command": "session.close", "args": {"save": True}},
        ],
        timeout=180,
    )
    if any(not item.get("success", False) for item in results):
        raise ExcelCliError(str(results))


def _formula(path: Path, sheet: str, cell: str) -> str:
    workbook = load_workbook(path, data_only=False)
    try:
        value = workbook[sheet][cell].value
        return "" if value is None else str(value)
    finally:
        workbook.close()


def _cached(path: Path, sheet: str, cell: str):
    workbook = load_workbook(path, data_only=True)
    try:
        return workbook[sheet][cell].value
    finally:
        workbook.close()


def test_style_screenshot_and_file_released(tmp_path: Path):
    _require_excel()
    path = tmp_path / "style.xlsx"
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active["A1"] = "x"
    workbook.save(path)
    workbook.close()
    body, shots = _apply(path, [{"action": "format", "sheet": "Sheet1", "range": "A1", "fontColor": "#FF0000"}])
    assert body["engine"] == "openpyxl"
    assert body["calc"] == "skipped"
    assert body["calculated"] is False
    assert body["committed"] is True
    assert body["verified"] is True
    assert body["ok"] is True
    assert shots and len(shots[0].data) > 100
    assert file_is_locked(str(path)) is False


def test_parallel_files_and_sequential_same_file(tmp_path: Path):
    _require_excel()
    first = tmp_path / "one.xlsx"
    second = tmp_path / "two.xlsx"
    _create(first)
    _create(second)
    results = {}

    def run(label, path):
        results[label] = _apply(path, [{"action": "set_values", "sheet": "Sheet1", "range": "A1", "values": [[label]]}])

    threads = [
        threading.Thread(target=run, args=("one", first)),
        threading.Thread(target=run, args=("two", second)),
    ]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(180)
    assert results["one"][0]["ok"] is True
    assert results["two"][0]["ok"] is True
    again, _ = _apply(first, [{"action": "set_values", "sheet": "Sheet1", "range": "B1", "values": [[2]]}])
    assert again["ok"] is True
    assert _formula(first, "Sheet1", "B1") == "2"


def test_read_returns_cached_value_after_input_change(tmp_path: Path):
    _require_excel()
    path = tmp_path / "calc.xlsx"
    _create(path)
    created, _ = _apply(
        path,
        [
            {"action": "create_sheet", "name": "汇总"},
            {"action": "set_values", "sheet": "Sheet1", "range": "A1", "values": [[2]]},
            {"action": "set_formulas", "sheet": "汇总", "range": "B1", "formulas": [["=Sheet1!A1"]]},
        ],
    )
    assert created["calculated"] is True, created
    changed, _ = _apply(path, [{"action": "set_values", "sheet": "Sheet1", "range": "A1", "values": [[9]]}])
    assert changed["calc"] == "done"
    assert changed["ok"] is True, changed
    assert _cached(path, "汇总", "B1") == 9
    page = read_workbook(str(path), sheet="汇总", range="B1", mode="dense")
    assert page["values"][0][0] == 9


def test_rename_rewrites_cross_sheet_formula_and_name(tmp_path: Path):
    _require_excel()
    path = tmp_path / "rename.xlsx"
    _create(path)
    setup, _ = _apply(
        path,
        [
            {"action": "create_sheet", "name": "汇总"},
            {"action": "set_values", "sheet": "Sheet1", "range": "A1", "values": [[10]]},
            {"action": "set_formulas", "sheet": "汇总", "range": "B1", "formulas": [["=Sheet1!A1"]]},
            {"action": "define_name", "name": "Source", "formula": "=Sheet1!$A$1"},
        ],
    )
    assert setup["ok"] is True, setup
    renamed, _ = _apply(path, [{"action": "rename_sheet", "oldName": "Sheet1", "newName": "数据"}])
    assert renamed["engine"] == "com"
    assert renamed["ok"] is True, renamed
    assert _formula(path, "汇总", "B1").replace("$", "") == "=数据!A1"
    workbook = load_workbook(path)
    try:
        assert "数据" in workbook.defined_names["Source"].attr_text
    finally:
        workbook.close()


def test_insert_keeps_cross_sheet_and_lambda_pointing_at_value(tmp_path: Path):
    _require_excel()
    path = tmp_path / "insert.xlsx"
    _create(path)
    setup, _ = _apply(
        path,
        [
            {"action": "create_sheet", "name": "汇总"},
            {"action": "set_values", "sheet": "Sheet1", "range": "A1", "values": [["n"], [10]]},
            {"action": "set_formulas", "sheet": "汇总", "range": "B1", "formulas": [["=Sheet1!A2"]]},
            {"action": "define_name", "name": "AddOne", "formula": "=LAMBDA(x,x+1)"},
            {"action": "set_formulas", "sheet": "汇总", "range": "C1", "formulas": [["=AddOne(Sheet1!A2)"]]},
        ],
    )
    assert setup["ok"] is True, setup
    before = file_hash(str(path))
    failed, _ = _apply(path, [{"action": "insert_rows", "sheet": "Missing", "row": 2, "count": 1}])
    assert failed["committed"] is False
    assert file_hash(str(path)) == before
    inserted, shots = _apply(path, [{"action": "insert_rows", "sheet": "Sheet1", "row": 2, "count": 1}])
    assert inserted["ok"] is True, inserted
    assert _formula(path, "汇总", "B1").replace("$", "") == "=Sheet1!A3"
    assert "A3" in _formula(path, "汇总", "C1").replace("$", "")
    workbook = load_workbook(path)
    try:
        assert "LAMBDA" in workbook.defined_names["AddOne"].attr_text.upper()
    finally:
        workbook.close()
    assert shots


def test_dynamic_array_spill_is_in_the_screenshot(tmp_path: Path):
    _require_excel()
    path = tmp_path / "spill.xlsx"
    _create(path)
    body, shots = _apply(path, [{"action": "set_formulas", "sheet": "Sheet1", "range": "A1", "formulas": [["=SEQUENCE(5)"]]}])
    assert body["ok"] is True, body
    assert body["screenshots"]
    address = body["screenshots"][0]["range"]
    assert address != "A1"
    assert shots[0].data


def test_verify_does_not_replay_insert(tmp_path: Path):
    _require_excel()
    path = tmp_path / "retry.xlsx"
    _create(path)
    inserted, _ = _apply(
        path,
        [
            {"action": "set_values", "sheet": "Sheet1", "range": "A1", "values": [["keep"], ["row"]]},
            {"action": "insert_rows", "sheet": "Sheet1", "row": 2, "count": 1},
        ],
    )
    assert inserted["committed"] is True, inserted
    assert _formula(path, "Sheet1", "A1") == "keep"
    assert _formula(path, "Sheet1", "A3") == "row"
    verified, _ = _apply(path, [{"action": "verify", "ranges": {"Sheet1": "A1:A3"}}])
    assert verified["ok"] is True, verified
    assert _formula(path, "Sheet1", "A3") == "row"
    assert _formula(path, "Sheet1", "A4") in {"", "None"}


def test_read_waits_for_write_lock(tmp_path: Path):
    _require_excel()
    from excel_sovereign.lock import FileLock

    path = tmp_path / "locked.xlsx"
    _create(path)
    holder = FileLock(str(path))
    holder.acquire()
    seen = {}

    def read():
        seen["page"] = read_workbook(str(path), sheet="Sheet1", range="A1")

    thread = threading.Thread(target=read)
    thread.start()
    time.sleep(0.4)
    assert thread.is_alive()
    holder.release()
    thread.join(15)
    assert not thread.is_alive()
    assert seen["page"]["ok"] is True


def _assert_commit_matches_bytes(path: Path, before: str, body: dict) -> None:
    after = file_hash(str(path))
    if body["committed"]:
        assert after != before
        assert body["saved"] is True
    else:
        assert body["error"]["code"] == "op_failed"
        assert after == before


def test_filtered_table_insert_commit_matches_file(tmp_path: Path):
    """Excel accepts EntireRow.Insert on this filtered table. A rejected insert still does not commit."""
    _require_excel()
    path = tmp_path / "filtered.xlsx"
    _create(path)
    setup, _ = _apply(
        path,
        [
            {"action": "set_values", "sheet": "Sheet1", "range": "A1", "values": [["n"], [1], [2]]},
            {
                "action": "table_create",
                "args": {"sheetName": "Sheet1", "tableName": "Data", "rangeAddress": "A1:A3", "hasHeaders": True},
            },
            {
                "action": "table_apply_filter",
                "args": {"tableName": "Data", "columnName": "n", "criteria": "=1"},
            },
        ],
        "excel_table",
    )
    assert setup["committed"] is True, setup
    before = file_hash(str(path))
    body, _ = _apply(path, [{"action": "insert_rows", "sheet": "Sheet1", "row": 3, "count": 1}])
    _assert_commit_matches_bytes(path, before, body)
    rejected_before = file_hash(str(path))
    rejected, _ = _apply(path, [{"action": "insert_rows", "sheet": "Missing", "row": 1, "count": 1}])
    assert rejected["committed"] is False
    assert file_hash(str(path)) == rejected_before


def test_merged_insert_commit_matches_file(tmp_path: Path):
    """Excel accepts EntireRow.Insert through this merge. A rejected insert still does not commit."""
    _require_excel()
    path = tmp_path / "merged.xlsx"
    _create(path)
    setup, _ = _apply(
        path,
        [
            {"action": "set_values", "sheet": "Sheet1", "range": "A1", "values": [["h"], ["a"], ["b"]]},
            {"action": "merge", "sheet": "Sheet1", "range": "A2:A3"},
            {"action": "create_sheet", "name": "汇总"},
            {"action": "set_formulas", "sheet": "汇总", "range": "B1", "formulas": [["=Sheet1!A1"]]},
        ],
    )
    assert setup["committed"] is True, setup
    before = file_hash(str(path))
    body, _ = _apply(path, [{"action": "insert_rows", "sheet": "Sheet1", "row": 2, "count": 1}])
    _assert_commit_matches_bytes(path, before, body)
    if body["committed"]:
        assert _formula(path, "汇总", "B1").replace("$", "") == "=Sheet1!A1"


def test_layout_finance_and_analytics(tmp_path: Path):
    _require_excel()
    from excel_sovereign.book import create_empty
    from test_layout import _analytics, _catalog, _finance

    finance = tmp_path / "finance.xlsx"
    create_empty(str(finance))
    body, shots = _apply(finance, [_finance(), _catalog()])
    assert body["engine"] == "openpyxl", body
    assert body["ok"] is True, body
    assert shots and shots[0].data
    workbook = load_workbook(finance)
    try:
        assert workbook["Sheet1"]["A1"].value == "收入预测"
        assert workbook["Sheet1"]["B4"].font.color.rgb[-6:] == "0000FF"
        assert workbook["Sheet1"]["B5"].font.color.rgb[-6:] == "008000"
        assert workbook["科目"]["A6"].border.top.style == "thin"
    finally:
        workbook.close()
    analytics = tmp_path / "analytics.xlsx"
    create_empty(str(analytics))
    stacked, images = _apply(analytics, [_analytics()])
    assert stacked["engine"] == "com", stacked
    assert stacked["committed"] is True, stacked
    assert images
    with zipfile.ZipFile(analytics) as archive:
        assert any(name.startswith("xl/tables/") for name in archive.namelist())


def test_hyperlink_hide_and_measure(tmp_path: Path):
    _require_excel()
    path = tmp_path / "more.xlsx"
    _create(path)
    table, _ = _apply(
        path,
        [
            {"action": "set_values", "sheet": "Sheet1", "range": "A1", "values": [["n", "v"], [1, 2], [3, 4]]},
            {
                "action": "table_create",
                "args": {"sheetName": "Sheet1", "tableName": "Data", "rangeAddress": "A1:B3", "hasHeaders": True},
            },
        ],
        "excel_table",
    )
    assert table["ok"] is True, table
    view, _ = _apply(
        path,
        [
            {
                "action": "hyperlink_add",
                "args": {
                    "sheetName": "Sheet1",
                    "cellAddress": "A1",
                    "url": "https://example.com",
                    "displayText": "n",
                },
            },
            {"action": "create_sheet", "name": "Side"},
            {"action": "sheet_hide", "sheet": "Side"},
            {"action": "sheet_show", "sheet": "Side"},
        ],
        "excel_view",
    )
    assert view["ok"] is True, view
    model, _ = _apply(
        path,
        [
            {"action": "table_add_to_data_model", "args": {"tableName": "Data"}},
            {
                "action": "datamodel_create_measure",
                "args": {
                    "tableName": "Data",
                    "measureName": "TotalV",
                    "daxFormula": "SUM(Data[v])",
                },
            },
        ],
        "excel_model",
    )
    assert model["ok"] is True, model


def test_protected_insert_is_not_committed(tmp_path: Path):
    _require_excel()
    path = tmp_path / "protected.xlsx"
    _create(path)
    workbook = load_workbook(path)
    workbook.active.protection.sheet = True
    workbook.save(path)
    workbook.close()
    before = file_hash(str(path))
    body, _ = _apply(path, [{"action": "insert_rows", "sheet": "Sheet1", "row": 1, "count": 1}])
    assert body["committed"] is False
    assert body["error"]["code"] == "op_failed"
    assert file_hash(str(path)) == before


def test_font_change_on_pivot_workbook_uses_com_and_keeps_pivot(tmp_path: Path):
    _require_excel()
    path = tmp_path / "pivot.xlsx"
    _create(path)
    setup, _ = _apply(
        path,
        [
            {"action": "set_values", "sheet": "Sheet1", "range": "A1", "values": [["group", "amount"], ["a", 1], ["a", 2]]},
            {
                "action": "pivot_create_from_range",
                "args": {
                    "sourceSheet": "Sheet1",
                    "sourceRange": "A1:B3",
                    "destinationSheet": "Sheet1",
                    "destinationCell": "D1",
                    "pivotTableName": "AmountByGroup",
                },
            },
        ],
        "excel_table",
    )
    assert setup["committed"] is True, setup
    info = inspect_workbook(str(path))
    assert info.has_com_parts
    before_parts = set(info.com_parts)
    tint, _ = _apply(path, [{"action": "format", "sheet": "Sheet1", "range": "A1", "fontColor": "#FF0000"}])
    assert tint["engine"] == "com", tint
    assert tint["committed"] is True, tint
    after = set(inspect_workbook(str(path)).com_parts)
    assert before_parts <= after


def test_xlsm_roundtrip_keeps_vba_part(tmp_path: Path):
    _require_excel()
    path = tmp_path / "macro.xlsm"
    _create(path, macro=True)
    imported, _ = _apply(
        path,
        [{"action": "vba_import", "module": "Mod1", "code": "Sub Ping()\nEnd Sub\n"}],
        "excel_vba",
    )
    assert imported["committed"] is True, imported
    with zipfile.ZipFile(path) as archive:
        before = archive.read("xl/vbaProject.bin")
    body, _ = _apply(path, [{"action": "format", "sheet": "Sheet1", "range": "A1", "fontColor": "#0000FF"}])
    assert body["committed"] is True, body
    with zipfile.ZipFile(path) as archive:
        assert archive.read("xl/vbaProject.bin") == before
    unchanged = file_hash(str(path))
    viewed, shots = _apply(path, [{"action": "vba_view", "module": "Mod1"}], "excel_vba")
    assert viewed["ok"] is True and viewed["readOnly"] is True, viewed
    assert viewed["committed"] is False and shots == []
    assert "Ping" in json.dumps(viewed["results"][0]["result"], ensure_ascii=False)
    assert file_hash(str(path)) == unchanged


def test_table_chart_validation_comment_and_query(tmp_path: Path):
    _require_excel()
    path = tmp_path / "phase4.xlsx"
    _create(path)
    table, _ = _apply(
        path,
        [
            {"action": "set_values", "sheet": "Sheet1", "range": "A1", "values": [["n", "v"], [1, 2], [3, 4]]},
            {
                "action": "table_create",
                "args": {"sheetName": "Sheet1", "tableName": "Data", "rangeAddress": "A1:B3", "hasHeaders": True},
            },
            {"action": "table_append", "args": {"tableName": "Data", "rows": [[5, 6]]}},
            {
                "action": "chart_create_from_range",
                "args": {
                    "sheetName": "Sheet1",
                    "sourceRangeAddress": "A1:B4",
                    "chartType": "ColumnClustered",
                    "targetRange": "D2:H12",
                },
            },
            {"action": "table_list"},
        ],
        "excel_table",
    )
    assert table["ok"] is True, table
    assert [item["action"] for item in table["results"]] == ["table_list"]
    assert "Data" in json.dumps(table["results"][0]["result"])
    view, _ = _apply(
        path,
        [
            {
                "action": "validation_add",
                "args": {
                    "sheetName": "Sheet1",
                    "rangeAddress": "A2:A10",
                    "validationType": "whole",
                    "validationOperator": "greaterThan",
                    "formula1": "0",
                },
            },
            {"action": "comment_set", "sheet": "Sheet1", "cell": "B2", "text": "note"},
            {
                "action": "conditional_format_add",
                "args": {
                    "sheetName": "Sheet1",
                    "rangeAddress": "B2:B10",
                    "ruleType": "cellValue",
                    "operatorType": "greaterThan",
                    "formula1": "0",
                    "interiorColor": "#C6EFCE",
                },
            },
            {"action": "freeze", "sheet": "Sheet1", "rows": 1, "columns": 0},
        ],
        "excel_view",
    )
    assert view["ok"] is True, view
    model, _ = _apply(
        path,
        [
            {
                "action": "powerquery_create",
                "args": {
                    "queryName": "Tiny",
                    "mCode": "let Source = #table({\"n\"}, {{1}}) in Source",
                    "loadDestination": "ConnectionOnly",
                },
            },
        ],
        "excel_model",
    )
    assert model["ok"] is True, model
    comment = read_workbook(str(path), sheet="Sheet1", range="A4")
    assert comment["ok"] is True


def test_batched_failure_reports_the_failing_op_and_keeps_bytes(tmp_path: Path):
    _require_excel()
    path = tmp_path / "batch.xlsx"
    _create(path)
    before = file_hash(str(path))
    body, _ = _apply(
        path,
        [
            {"action": "insert_rows", "sheet": "Sheet1", "row": 1, "count": 1},
            {"action": "delete_sheet", "sheet": "Missing"},
            {"action": "insert_rows", "sheet": "Sheet1", "row": 1, "count": 1},
        ],
    )
    assert body["committed"] is False
    assert body["failedAt"] == 1, body
    assert body["completedOps"] == 1
    assert file_hash(str(path)) == before


def test_trim_sheet_drops_formatting_past_the_data(tmp_path: Path):
    _require_excel()
    path = tmp_path / "trim.xlsx"
    _create(path)
    workbook = load_workbook(path)
    ws = workbook["Sheet1"]
    ws["A1"] = "keep"
    ws["B3"] = "=A1"
    for row in range(1, 60):
        ws.cell(row, 40).font = Font(bold=True)
    ws.column_dimensions["AZ"].width = 25
    workbook.save(path)
    workbook.close()
    assert read_workbook(str(path), mode="overview")["sheets"][0]["extent"] == "A1:AZ59"
    body, _ = _apply(path, [{"action": "trim_sheet", "sheet": "Sheet1"}])
    assert body["committed"] is True, body
    sheet = read_workbook(str(path), mode="overview")["sheets"][0]
    assert sheet["usedRange"] == "A1:B3"
    assert "extent" not in sheet, sheet
    assert _formula(path, "Sheet1", "B3") == "=A1"
